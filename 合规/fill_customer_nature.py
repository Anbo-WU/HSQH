from pathlib import Path
import openpyxl


def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip().replace(" ", "").lower()


def find_first_data_row(ws):
    # 兼容两种情况：
    # 1) 第一行就是数据行（如 A.xlsx）
    # 2) 第一行是表头（如 B.xlsx）
    first_row = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    header_keywords = {"编号", "客户名称", "企业名称", "客户性质", "名称", "序号"}
    if any(isinstance(v, str) and v in header_keywords for v in first_row):
        return 2
    return 1


def load_workbook(path):
    return openpyxl.load_workbook(path, data_only=False)


def build_name_to_nature_map(ws, name_col_letter="B", nature_col_letter="D"):
    name_idx = openpyxl.utils.column_index_from_string(name_col_letter) - 1
    nature_idx = openpyxl.utils.column_index_from_string(nature_col_letter) - 1

    start_row = find_first_data_row(ws)
    mapping = {}

    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if len(row) <= max(name_idx, nature_idx):
            continue

        name = row[name_idx]
        nature = row[nature_idx]
        if name is None or str(name).strip() == "":
            continue

        mapping[normalize_text(name)] = nature

    return mapping


def fill_customer_nature(source_path, reference_path, output_path=None):
    source_wb = load_workbook(source_path)
    reference_wb = load_workbook(reference_path)

    source_ws = source_wb.active
    reference_ws = reference_wb.active

    mapping = build_name_to_nature_map(reference_ws)
    start_row = find_first_data_row(source_ws)

    updated_count = 0
    not_found_count = 0

    for row_idx in range(start_row, source_ws.max_row + 1):
        name_cell = source_ws.cell(row_idx, 2)  # B列
        target_cell = source_ws.cell(row_idx, 3)  # C列

        name = name_cell.value
        if name is None or str(name).strip() == "":
            continue

        normalized_name = normalize_text(name)
        if normalized_name in mapping:
            matched_value = mapping[normalized_name]
            if target_cell.value in (None, ""):
                target_cell.value = matched_value
                updated_count += 1
            else:
                # 如果 C 列已有值，则保留现有值，不覆盖
                pass
        else:
            not_found_count += 1

    if output_path is None:
        output_path = source_path

    source_wb.save(output_path)
    return updated_count, not_found_count


if __name__ == "__main__":
    base_dir = Path(__file__).resolve().parent
    workspace_dir = base_dir.parent

    candidates = [
        base_dir / "A.xlsx",
        workspace_dir / "A.xlsx",
        Path("A.xlsx"),
    ]
    source_path = None
    for candidate in candidates:
        if candidate.exists():
            source_path = candidate
            break

    reference_candidates = [
        base_dir / "B.xlsx",
        workspace_dir / "B.xlsx",
        Path("B.xlsx"),
    ]
    reference_path = None
    for candidate in reference_candidates:
        if candidate.exists():
            reference_path = candidate
            break

    if source_path is None or reference_path is None:
        raise FileNotFoundError("未找到 A.xlsx 或 B.xlsx")

    updated_count, not_found_count = fill_customer_nature(source_path, reference_path)
    print(f"处理完成：已填充 {updated_count} 条，未匹配 {not_found_count} 条")
    print(f"输出文件：{source_path}")
