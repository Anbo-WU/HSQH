#!/usr/bin/env python3
"""从客户开户材料批量生成两张客户信息登记表；所有 PDF 均在本地识别。"""
from __future__ import annotations

import argparse
from copy import copy
from dataclasses import asdict, dataclass, replace
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
import pymupdf

from ocr_support import Line, PDFReader, rows

BASE = Path(__file__).resolve().parent
SHEET1 = '场外衍生品协议的签署'
SHEET2 = '交易人员及信息'
KYC_MAP = {
    'B': ('基本信息表', 'C6'), 'F': ('交易信息', 'C8'),
    'G': ('交易信息', 'C7'), 'H': ('交易信息', 'C10'),
    'I': ('基本信息表', 'C12'), 'J': ('基本信息表', 'C13'),
    'M': ('基本信息表', 'C7'), 'N': ('基本信息表', 'C14'),
    'O': ('基本信息表', 'C21'),
    'P': ('基本信息表', 'C23'), 'Q': ('基本信息表', 'C24'),
    'U': ('交易信息', 'C9'), 'V': ('交易信息', 'C10'),
    'W': ('交易信息', 'C11'), 'X': ('交易信息', 'C12'),
    'Y': ('交易信息', 'C15'), 'Z': ('交易信息', 'C17'),
    'AA': ('交易信息', 'C16'), 'AB': ('交易信息', 'C18'),
    'AI': ('基本信息表', 'D52'), 'AJ': ('基本信息表', 'D61'),
}
AUTH_MAP = {'C': 'B14', 'E': 'H14', 'F': 'E14', 'G': 'F14', 'H': 'I14'}
DATE_RE = re.compile(r'(?<!\d)((?:19|20)\d{2})\s*[年./-]\s*(\d{1,2})\s*[月./-]\s*(\d{1,2})(?:日)?(?!\d)')


def compact(value) -> str:
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(value or '')))


def dates(text: str) -> list[str]:
    result = []
    for match in DATE_RE.finditer(unicodedata.normalize('NFKC', text)):
        try:
            result.append(date(*map(int, match.groups())).isoformat())
        except ValueError:
            continue
    return result


def expiry(value) -> str:
    text = compact(value)
    if text.endswith(('长期', '永久', '长期有效')):
        return '长期'
    found = list(DATE_RE.finditer(text))
    if not found or found[-1].end() != len(text):
        raise ValueError('未识别出有效截止日期')
    try:
        end = date(*map(int, found[-1].groups()))
    except ValueError as exc:
        raise ValueError('证件截止日期不合法') from exc
    return end.isoformat().replace('-', '.')


def chinese_number(text: str) -> Decimal:
    text = text.translate(str.maketrans('零〇一二两三四五六七八九壹贰叁肆伍陆柒捌玖拾佰仟萬億',
                                       '001223456789123456789十百千万亿'))
    for unit, amount in [('亿', 100000000), ('万', 10000)]:
        if unit in text:
            left, right = text.split(unit, 1)
            return chinese_number(left or '1') * amount + (chinese_number(right) if right else 0)
    if re.fullmatch(r'\d+(?:\.\d+)?', text):
        return Decimal(text)
    total = Decimal(0)
    digits = ''
    for char in text:
        if char.isdigit():
            digits += char
        elif char in '十百千':
            total += Decimal(digits or '1') * {'十': 10, '百': 100, '千': 1000}[char]
            digits = ''
        else:
            raise ValueError(f'无法解析中文金额：{text}')
    return total + Decimal(digits or '0')


def capital_wan(text: str) -> float:
    text = compact(text).replace(',', '')
    if any(currency in text for currency in ('美元', '港元', '欧元', '日元')):
        raise ValueError('注册资本为外币，不能直接填入人民币万元')
    text = re.sub(r'^.*?注册资本[:：]?', '', text)
    text = text.replace('人民币', '').replace('整', '').replace('圆', '元')
    match = re.fullmatch(r'([\d.零〇一二两三四五六七八九壹贰叁肆伍陆柒捌玖十百千拾佰仟万亿萬億]+)元?', text)
    if not match or ('元' not in text and not any(x in text for x in '万亿萬億')):
        raise ValueError('注册资本数值或单位不明确')
    return float(chinese_number(match[1]) / Decimal(10000))


def credit_class(score: float) -> str:
    if score >= 80:
        return '正常'
    if 60 < score < 80:
        return '关注'
    raise ValueError('分数等于或低于 60，用户尚未定义资信分类规则')


def suitability(folder_name: str) -> str:
    match = re.search(r'(C4低买高|C4|C5|B类专业交易者)$', compact(folder_name), re.I)
    return match[1].upper() if match and match[1].upper() in {'C4', 'C5'} else match[1] if match else '专业交易者'


def row_text(group: list[Line]) -> str:
    return ' '.join(line.text for line in group)


def label_value(lines: list[Line], label: str) -> tuple[str, list[Line]]:
    for group in rows(lines):
        for i, line in enumerate(group):
            text = compact(line.text)
            if label in text:
                tail = text.split(label, 1)[1].lstrip(':：')
                chosen = group[i+1:]
                value = tail or ''.join(x.text for x in chosen).strip()
                if value:
                    return value, [line] + chosen
    raise ValueError(f'没有找到“{label}”对应内容')


def beneficiary_names(pages: list[list[Line]]) -> tuple[str, list[Line]]:
    names = []
    evidence = []
    uncertain = []
    for lines in pages:
        groups = rows(lines)
        for index, group in enumerate(groups):
            context = compact(' '.join(row_text(r) for r in groups[max(0, index-2):index+1]))
            if not any(word in context for word in ('股东/控制人/高管', '股东、控制人、高管')):
                continue
            for i, line in enumerate(group):
                if not re.match(r'^\*?姓名[:：]?', compact(line.text)):
                    continue
                inline = re.sub(r'^\*?姓名[:：]?', '', compact(line.text))
                values = []
                for item in group[i+1:]:
                    if any(word in compact(item.text) for word in ('证件', '国籍', '地址')):
                        break
                    values.append(item)
                name = inline or ''.join(compact(item.text) for item in values)
                if name and re.fullmatch(r'[\u4e00-\u9fffA-Za-z·.]{2,80}', name):
                    if name not in names:
                        names.append(name)
                    evidence.extend([line] + values)
                elif name or re.search(r'\d{15,18}[Xx]?', row_text(group)):
                    uncertain.append(row_text(group))
    if uncertain:
        raise ValueError('部分受益人姓名未识别完整；已识别：' + '、'.join(names) + '；请核对整份采集表')
    if not names:
        raise ValueError('没有识别出“股东/控制人/高管”姓名')
    return '、'.join(names), evidence


def money(text: str) -> Decimal | None:
    text = compact(text).replace(',', '').replace('−', '-').replace('—', '-')
    if re.fullmatch(r'\(\d+\.\d{2}\)', text):
        text = '-' + text[1:-1]
    # 行次整数不能冒充金额；没有小数的金额也支持，但由调用方的列位置排除行次。
    if not re.fullmatch(r'-?\d+(?:\.\d{1,2})?', text):
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


FINANCE_HEADERS = {'行次', '期末余额', '期末数', '年末余额', '年末数',
                   '年初余额', '年初数', '期初余额', '期初数', '本期金额',
                   '本年累计金额', '本年累计数', '本年金额', '上年金额',
                   '上期金额', '上年累计金额', '本月金额', '本月数'}
FINANCE_TITLES = ('资产负债表', '利润表', '损益表', '现金流量表')
FINANCE_LABELS = {'资产合计': {'资产合计', '资产总计', '总资产'},
                  '负债合计': {'负债合计', '负债总计', '总负债'},
                  '净利润': {'净利润'}}
HEADER_ALIASES = {'期末余额': {'期末余额', '期末数', '年末余额', '年末数'},
                  '本年累计金额': {'本年累计金额', '本年累计数', '本年金额'}}


def is_financial_label(text: str, label: str) -> bool:
    text = compact(text)
    # 相邻左格金额可能被 OCR 连到右格行标题，不把金额当作标题的一部分。
    text = re.sub(r'^-?[\d,]+\.\d{2}(?=[\u4e00-\u9fff])', '', text)
    text = re.sub(r'^(?:[一二三四五六七八九十]+|\d+)[、.．]', '', text)
    text = re.split(r'[（(]', text, 1)[0].rstrip(':：')
    return text in FINANCE_LABELS[label]


def financial_value(lines: list[Line], label: str, header: str) -> tuple[float, list[Line]]:
    candidates = []
    for group in rows(lines):
        for line in group:
            if not is_financial_label(line.text, label):
                continue
            # 同一行出现另一组科目时，不能跨过它去取另一半表格的金额。
            boundary = min((x.x0 for x in group if x.x0 > line.x0
                            and re.search(r'[\u4e00-\u9fff]', x.text)), default=1.0)
            headings = [x for x in lines if compact(x.text) in HEADER_ALIASES.get(header, {header})
                        and x.cy < line.cy and line.x0 < x.x0 < boundary]
            if not headings:
                continue
            heading = min(headings, key=lambda x: (line.cy-x.cy, x.x0))
            peer_headers = [x for x in lines if abs(x.cy-heading.cy) < 0.01 and
                            compact(x.text) in FINANCE_HEADERS]
            center = (heading.x0 + heading.x1)/2
            left = max(((x.x0+x.x1)/2 for x in peer_headers if (x.x0+x.x1)/2 < center), default=line.x1)
            right = min(((x.x0+x.x1)/2 for x in peer_headers if (x.x0+x.x1)/2 > center), default=2*boundary-center)
            values = [x for x in group if (left+center)/2 < (x.x0+x.x1)/2 < (center+right)/2 and money(x.text) is not None]
            if len(values) == 1:
                candidates.append((money(values[0].text), [line, heading, values[0]]))
    if len(candidates) != 1:
        raise ValueError(f'无法唯一定位“{label} / {header}”金额')
    value, evidence = candidates[0]
    unit_lines = [x for x in lines if '单位' in compact(x.text)]
    unit_text = ''.join(compact(x.text) for x in unit_lines)
    units = set(re.findall(r'单位[:：]?(?:人民币)?(万元|千元|元)', unit_text))
    if len(units) != 1:
        raise ValueError('财报计量单位未识别，不能确认金额为元')
    value *= {'元': 1, '千元': 1000, '万元': 10000}[units.pop()]
    return float(value), evidence + unit_lines


def continuation_headers(previous: list[Line], current: list[Line]) -> list[Line]:
    """仅在无新表头、行次列对齐且连续时，沿用紧邻上一页的表头和单位。"""
    if any(compact(x.text) in FINANCE_HEADERS or any(t in compact(x.text) for t in FINANCE_TITLES)
           for x in current):
        return []
    row_headers = [x for x in previous if compact(x.text) == '行次']
    if not row_headers:
        return []
    for heading in row_headers:
        center = (heading.x0+heading.x1)/2
        def row_numbers(items):
            return [int(compact(x.text)) for x in items if x.cy > heading.cy
                    and abs((x.x0+x.x1)/2-center) < max((heading.x1-heading.x0)/2, .012)
                    and x.score >= .85 and re.fullmatch(r'\d{1,3}', compact(x.text))]
        before = row_numbers(previous)
        # 当前页的纵坐标重新起算，只比较横向行次列。
        after = row_numbers([replace(x, y0=x.y0+1, y1=x.y1+1) for x in current])
        if not before or not after or min(after) != max(before)+1:
            return []
    return [replace(x, y0=x.y0-1, y1=x.y1-1) for x in previous
            if compact(x.text) in FINANCE_HEADERS or '单位' in compact(x.text)]


def question18(lines: list[Line]) -> str:
    paragraph = []
    active = False
    for group in rows(lines):
        text = compact(row_text(group))
        if re.match(r'^(?:第)?18[.、:：题]', text):
            active = True
        elif active and re.match(r'^(?:[A-G][.、:：]|19[.、:：])', text, re.I):
            break
        if active:
            paragraph.append(text)
    match = re.search(r'(?<![A-Za-z])([A-D])(?:[)）])?\s*$', ''.join(paragraph), re.I)
    if not match:
        raise ValueError('第 18 题题干末尾未识别出 A–D；勾选式答案按要求不处理')
    return match[1].upper()


@dataclass
class Field:
    value: str | int | float | None = None
    source: str = ''
    evidence: str = ''
    status: str = '已提取'
    note: str = ''


def extracted(value, source: str, lines: list[Line], note='') -> Field:
    if any(line.score < 0.85 for line in lines):
        return Field(None, source, row_text(lines), '待核对', f'OCR 置信度不足；候选值：{value}。{note}')
    return Field(value, source, row_text(lines), note=note)


def credit_score(reader: PDFReader, path: Path, index: int, lines: list[Line]) -> Field:
    source = f'{path.name} / 第{index+1}页各处综合评分'
    scores, evidence, notes, errors = [], [], [], []

    def parse(text):
        text = re.sub(r'^综合评分[:：]?', '', compact(text)).removesuffix('分')
        return float(text) if re.fullmatch(r'-?\d{1,3}(?:\.\d+)?', text) and -100 <= float(text) <= 105 else None

    for group in rows(lines):
        for i, label in enumerate(group):
            # 只匹配填写栏；排除“1.综合评分60…”等分类说明，支持分数与标签连在一起。
            if not re.match(r'^综合评分(?:[:：]|-?\d|$)', compact(label.text)):
                continue
            tail = re.sub(r'^综合评分[:：]?', '', compact(label.text))
            candidates = [label] if tail else []
            if not tail:
                for item in group[i+1:]:
                    if re.search(r'[\u4e00-\u9fff]', item.text):
                        break
                    candidates.append(item)
            if len(candidates) != 1:
                errors.append('某处综合评分未能唯一定位数字')
                continue
            item = candidates[0]
            crop = (max(0, item.x0-.008), max(0, item.y0-.002),
                    min(1, item.x1+.008), min(1, item.y1+.002))
            verified, local_evidence = [], []
            for scale in (3.5, 5.0):
                # 正向填写栏中的手写数字可能被方向分类器误翻转，局部复核锁定原方向。
                retry = reader.page(path, index, crop=crop, scale=scale, use_cls=False)
                found = [(parse(x.text), x) for x in retry if parse(x.text) is not None]
                if len(found) == 1 and found[0][1].score >= .85:
                    verified.append(found[0][0])
                    local_evidence.append(found[0][1])
            if len(verified) != 2 or len(set(verified)) != 1:
                errors.append(f'局部复核未通过（原文：{item.text}；复核：{verified}）')
                continue
            scores.append(verified[0])
            evidence.extend([label]+local_evidence)
            notes.append(f'原文 {item.text} → 固定方向两种倍率复核 {verified[0]:g}')
    if errors or not scores or len(set(scores)) != 1:
        return Field(source=source, evidence=row_text(evidence), status='待核对',
                     note='；'.join(errors+notes+[f'综合评分未识别或多处不一致：{scores}']))
    return extracted(scores[0], source, evidence,
                     '；'.join(notes+[f'{len(scores)} 处综合评分复核一致']))


def unique_file(folder: Path, predicate) -> Path:
    found = sorted(p for p in folder.rglob('*') if p.is_file() and not p.name.startswith('~$') and predicate(p))
    if len(found) != 1:
        raise ValueError(f'{folder.name}：匹配文件数量为 {len(found)}，需要唯一文件')
    return found[0]


def document(folder: Path, number: int) -> Path:
    return unique_file(folder/'1.开户材料存档', lambda p: p.suffix.lower() == '.pdf' and re.match(rf'^{number}[.．、 _-]', p.name))


def get_sheet(book: Workbook, title: str) -> Worksheet:
    matches = [s for s in book if compact(s.title).removeprefix('sheet1').removeprefix('sheet2').removeprefix('sheet3') == title or compact(s.title).endswith(title)]
    if not matches and title == '基本信息表':
        matches = [s for s in book if compact(s.title).lower() == 'sheet1']
    if len(matches) != 1 or not isinstance(matches[0], Worksheet):
        raise ValueError(f'KYC 中无法唯一找到工作表：{title}')
    return matches[0]


def read_kyc(folder: Path) -> tuple[dict[str, Field], dict[str, Field]]:
    main, auth = {}, {}
    mapping = [(main, col, sheet, cell) for col, (sheet, cell) in KYC_MAP.items()]
    mapping += [(auth, col, '交易授权', cell) for col, cell in AUTH_MAP.items()]
    try:
        path = unique_file(folder/'0.OA', lambda p: 'KYC' in p.name.upper() and p.suffix.lower() == '.xlsx')
        book = load_workbook(path, data_only=True)
    except Exception as exc:
        for target, col, sheet, cell in mapping:
            target[col] = Field(source=f'KYC / {sheet} / {cell}', status='待核对', note=str(exc))
        return main, auth
    try:
        for target, col, sheet, cell in mapping:
            source = f'{path.relative_to(folder)} / {sheet} / {cell}'
            try:
                worksheet = get_sheet(book, sheet)
                source = f'{path.relative_to(folder)} / {worksheet.title} / {cell}'
                if target is auth and col == 'C':
                    source = f'{path.relative_to(folder)} / {worksheet.title} / B14:B{max(14, worksheet.max_row)}'
                    names, evidence = [], []
                    for row in worksheet.iter_rows(min_row=14, min_col=2, max_col=2):
                        item = row[0]
                        if item.data_type == 'e':
                            raise ValueError(f'来源单元格 {item.coordinate} 含错误')
                        if item.value is not None and str(item.value).strip():
                            names.append(str(item.value).strip())
                            evidence.append(f'{item.coordinate}：{item.value}')
                    if not names:
                        raise ValueError('B14 及以下未找到非空交易授权人')
                    target[col] = Field('，'.join(names), source, '\n'.join(evidence))
                    continue
                item = worksheet[cell]
                value = item.value
                if value is None or str(value).strip() == '' or item.data_type == 'e':
                    raise ValueError('来源单元格为空、公式无缓存或含错误')
                if isinstance(value, (date, datetime)):
                    value = value.strftime('%Y-%m-%d')
                elif isinstance(value, (float, int)):
                    if abs(value) >= 10**15:
                        raise ValueError('长号码以 Excel 数值保存，可能已丢失精度，请将源号码改为文本')
                    value = str(int(value)) if float(value).is_integer() else str(value)
                    if re.fullmatch(r'0+', item.number_format):
                        value = value.zfill(len(item.number_format))
                else:
                    value = str(value).strip()
                if target is main and col == 'Q':
                    value = expiry(value)
                target[col] = Field(value, source, str(item.value))
            except Exception as exc:
                target[col] = Field(source=source, status='待核对', note=str(exc))
    finally:
        book.close()
    return main, auth


def collect_customer(folder: Path, reader: PDFReader) -> tuple[dict, dict]:
    main, auth = read_kyc(folder)
    main['BA'] = Field(suitability(folder.name), '客户文件夹名称', folder.name)

    def task(columns, source, func):
        try:
            values = func()
            main.update(values)
        except Exception as exc:
            for column in columns:
                main[column] = Field(source=source, status='待核对', note=str(exc))

    def nature():
        path = document(folder, 1)
        value, evidence = label_value(reader.page(path, 0), '单位性质、资质')
        return {'K': extracted(value, f'{path.name} / 第1页 / 单位性质、资质', evidence)}

    def beneficiaries():
        path = document(folder, 13)
        with pymupdf.open(path) as doc:
            pages = [reader.page(path, i) for i in range(len(doc))]
        value, evidence = beneficiary_names(pages)
        return {'R': extracted(value, f'{path.name} / 全部页面 / 股东/控制人/高管', evidence)}

    def license_fields():
        path = document(folder, 14)
        lines = reader.page(path, 0)
        # 横置扫描件在页面中显示为纵向文字；尝试旋转后以完整识别文字量选正向。
        tall = sum(len(x.text) for x in lines if (x.y1-x.y0) > (x.x1-x.x0)*1.5)
        if tall > sum(len(x.text) for x in lines)*0.25:
            alternatives = [lines, reader.page(path, 0, rotation=90), reader.page(path, 0, rotation=270)]
            lines = max(alternatives, key=lambda ls: sum(len(x.text)*x.score for x in ls if x.x1-x.x0 > x.y1-x.y0))
        result = {}
        try:
            value, evidence = label_value(lines, '注册资本')
            # 同行可能还有其他栏目，仅保留从数额到单位的连续片段。
            amount = re.match(r'(?:人民币)?[\d,.零〇一二两三四五六七八九壹贰叁肆伍陆柒捌玖十百千拾佰仟万亿萬億]+(?:元|圆)(?:整)?', compact(value))
            if not amount:
                raise ValueError('未识别出注册资本完整金额及单位')
            result['AC'] = extracted(capital_wan(amount[0]), f'{path.name} / 注册资本（转换为万元）', evidence)
        except Exception as exc:
            result['AC'] = Field(source=path.name, status='待核对', note=str(exc))
        return result

    def finance():
        path = document(folder, 16)
        pages = []
        with pymupdf.open(path) as doc:
            for index in range(len(doc)):
                lines, rotation = reader.oriented_page(
                    path, index, FINANCE_TITLES + ('资产合计', '资产总计', '负债合计', '净利润'))
                # 印章可能挡住一侧表头；补读表头，金额仍取原始彩色扫描结果。
                if sum(compact(x.text) in HEADER_ALIASES['期末余额'] for x in lines) < sum(
                        compact(x.text) in {'年初余额', '年初数', '期初余额', '期初数'} for x in lines):
                    clean = reader.page(path, index, rotation=rotation, use_cls=False, suppress_red=True)
                    for item in clean:
                        if compact(item.text) not in FINANCE_HEADERS or item.score < .85:
                            continue
                        if not any(compact(x.text) == compact(item.text) and
                                   abs(x.x0-item.x0) < .025 and abs(x.cy-item.cy) < .025 for x in lines):
                            lines.append(item)
                inherited = continuation_headers(pages[-1][0], lines) if pages else []
                pages.append((lines+inherited, rotation, bool(inherited)))
        result = {}
        for col, label, heading in [('AD', '资产合计', '期末余额'),
                                    ('AE', '负债合计', '期末余额'),
                                    ('AG', '净利润', '本年累计金额')]:
            candidates, errors = [], []
            for index, (lines, rotation, inherited) in enumerate(pages):
                if not any(is_financial_label(x.text, label) for x in lines):
                    continue
                source = f'{path.name} / 第{index+1}页 / {label} / {heading} / 旋转{rotation}°'
                if inherited:
                    source += ' / 表头、单位沿用前页（行次连续且列对齐）'
                try:
                    value, evidence = financial_value(lines, label, heading)
                    candidates.append(extracted(value, source, evidence))
                except ValueError as exc:
                    errors.append(f'第{index+1}页：{exc}')
            if len(candidates) == 1 and not errors:
                result[col] = candidates[0]
            else:
                result[col] = Field(source=f'{path.name} / 全部页面 / {label} / {heading}',
                                    status='待核对', note='；'.join(errors) or
                                    f'关键词未识别或金额不能唯一确定（候选数量：{len(candidates)}）')
        return result

    def rating_date():
        path = document(folder, 5)
        lines = reader.page(path, 0)
        date_lines = [x for x in lines if '日期' in x.text]
        # 完整页漏掉手写年份的某个数字时，局部放大重试，而非猜测补年份。
        for item in date_lines[:]:
            if not dates(item.text):
                date_lines.remove(item)
                date_lines.extend(reader.page(path, 0, crop=(max(0,item.x0-.03), max(0,item.y0-.01), min(1,item.x1+.08), min(1,item.y1+.01)), scale=3.5))
        parsed = [(d, x) for x in date_lines for d in dates(x.text)]
        unique = set(d for d, _ in parsed)
        if len(unique) != 1:
            raise ValueError('第一页日期未识别或不同位置日期不一致')
        return {'AL': extracted(unique.pop(), f'{path.name} / 第1页日期', [x for _, x in parsed])}

    def loss_answer():
        path = document(folder, 2)
        lines = reader.page(path, 3)
        answer = question18(lines)
        # 只用题干末尾字母定位选项，不把选项前的勾号视为答案。
        option_lines = reader.page(path, 4)
        options = []
        for group in rows(option_lines):
            text = compact(row_text(group))
            if re.match(r'19[.、:：]', text):
                break
            match = re.match(rf'{answer}[.、:：](.+)', text)
            if match:
                options.append((match[1], group))
        if len(options) != 1 or '%' not in options[0][0]:
            raise ValueError(f'第18题答案为 {answer}，但未识别出该选项的损失比例')
        start = next(x.cy for x in lines if re.match(r'^(?:第)?18[.、:：题]', compact(x.text)))
        stem_lines = [x for x in lines if x.cy >= start-.005]
        return {'AM': extracted(options[0][0], f'{path.name} / 第4页第18题题干末尾及第5页选项', stem_lines+options[0][1], f'答案字母：{answer}')}

    def credit():
        path = unique_file(folder/'0.OA', lambda p: p.suffix.lower() == '.pdf' and '资信评估' in p.name)
        first, second = reader.page(path, 0), reader.page(path, 1)
        result = {}
        first_dates = [x for x in first if x.x0 > .45 and x.cy < .4 and ('日期' in x.text or dates(x.text))]
        second_dates = [x for x in second if x.x0 < .65 and x.cy > .55 and ('日期' in x.text or dates(x.text))]
        d1 = {d for x in first_dates for d in dates(x.text)}
        d2 = {d for x in second_dates for d in dates(x.text)}
        if len(d1) == 1 and d1 == d2:
            result['BB'] = extracted(next(iter(d1)), f'{path.name} / 第1页右上与第2页左下日期交叉校对', first_dates+second_dates)
        else:
            result['BB'] = Field(source=path.name, status='待核对', note=f'手写日期交叉校对未通过：第1页{sorted(d1)}；第2页{sorted(d2)}')
        try:
            result['BC'] = credit_score(reader, path, 1, second)
        except Exception as exc:
            result['BC'] = Field(source=path.name, status='待核对', note=f'综合评分复核失败：{exc}')
        return result

    task(['K'], '1.交易者基本信息表', nature)
    task(['R'], '13.受益人信息采集表', beneficiaries)
    task(['AC'], '14.营业执照', license_fields)
    task(['AD', 'AE', 'AG'], '16.财报表 / 全部页面关键词识别', finance)
    task(['AL'], '5.普通交易者适当性匹配意见告知书 / 第1页', rating_date)
    task(['AM'], '2.风险承受能力评估问卷 / 第4页第18题', loss_answer)
    task(['BB', 'BC'], '0.OA / 资信评估表', credit)
    if main['AD'].value is not None and main['AE'].value is not None:
        net = Decimal(str(main['AD'].value)) - Decimal(str(main['AE'].value))
        main['AF'] = Field(float(net), 'AD - AE', f'{main["AD"].value} - {main["AE"].value}', '已计算')
    else:
        main['AF'] = Field(source='AD - AE', status='待核对', note='总资产或总负债未确认')
    try:
        score = main['BC'].value
        if not isinstance(score, (int, float)):
            raise ValueError('资信评分未确认')
        main['AR'] = Field(credit_class(score), 'BC 分数分类', str(score), '已计算')
    except ValueError as exc:
        main['AR'] = Field(source='BC 分数分类', status='待核对', note=str(exc))
    return main, auth


def writable_cell(sheet: Worksheet, row: int, column: int) -> Cell:
    cell = sheet.cell(row, column)
    if not isinstance(cell, Cell):
        raise ValueError(f'模板 {sheet.title}!{cell.coordinate} 是合并单元格的非左上角位置，无法写入')
    return cell


def write_result(
    template: Path,
    output: Path,
    customers: list[tuple[Path, dict[str, Field], dict[str, Field]]],
) -> list[dict]:
    book = load_workbook(template)
    if SHEET1 not in book.sheetnames or SHEET2 not in book.sheetnames:
        raise ValueError('模板缺少指定的两张工作表')
    audit = []
    for sheet in (get_sheet(book, SHEET1), get_sheet(book, SHEET2)):
        defaults = [copy(writable_cell(sheet, 3, col)) for col in range(1, sheet.max_column+1)]
        for row in sheet.iter_rows(min_row=3):
            for cell in row:
                if isinstance(cell, Cell):
                    cell.value = None
        for rownum, (folder, main, auth) in enumerate(customers, 3):
            fields = main if sheet.title == SHEET1 else auth
            for colnum, original in enumerate(defaults, 1):
                cell = writable_cell(sheet, rownum, colnum)
                cell._style = copy(original._style)
                if sheet.title == SHEET1 and original.column_letter in {'AN','AO','AP','AQ','AS','AT','AU'}:
                    cell.value = original.value
            # sheet2 客户名用于关联；编号沿用客户根目录已有编号。
            match = re.match(r'^(\d+)[、.． _-]', folder.name)
            writable_cell(sheet, rownum, 1).value = match[1] if match else None
            if sheet.title == SHEET2:
                writable_cell(sheet, rownum, 2).value = main['B'].value
            for col, field in fields.items():
                cell = writable_cell(sheet, rownum, column_index_from_string(col))
                cell.value = field.value
                if isinstance(field.value, str):
                    cell.data_type = 's'
                    cell.number_format = '@'
                elif isinstance(field.value, (int, float)):
                    cell.number_format = '0.00' if col in {'AC','AD','AE','AF','AG'} else '0.##'
                cell.comment = Comment(f'来源：{field.source}\n状态：{field.status}\n识别原文：{field.evidence}\n{field.note}', '客户信息登记')
                if field.status == '待核对':
                    cell.fill = PatternFill('solid', fgColor='FFF2CC')
                audit.append({'客户文件夹': folder.name, '工作表': sheet.title,
                              '单元格': cell.coordinate, '字段': sheet[f'{col}2'].value, **asdict(field)})
            if sheet.title == SHEET1:
                sheet[f'AH{rownum}'] = f'=IF(AND(ISNUMBER(AD{rownum}),ISNUMBER(AE{rownum})),IFERROR(AE{rownum}/AD{rownum}*100,""),"")'
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    book.close()
    return audit


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=BASE, help='包含各客户大文件夹的根目录')
    parser.add_argument('--customer', help='只处理此客户文件夹名称')
    parser.add_argument('--template', type=Path, default=BASE/'场外衍生品中心客户信息登记表-模板.xlsx')
    parser.add_argument('--output', type=Path, help='输出 xlsx 路径，默认在结果目录生成时间戳文件')
    args = parser.parse_args()
    root = args.root.resolve()
    folders = sorted(p for p in root.iterdir() if p.is_dir() and ((p/'0.OA').exists() or (p/'1.开户材料存档').exists()))
    if args.customer:
        folders = [p for p in folders if p.name == args.customer]
    if not folders:
        parser.error('没有找到客户文件夹（应包含 0.OA 或 1.开户材料存档）')
    output = args.output or BASE/'结果'/f'客户信息登记表_{datetime.now():%Y%m%d_%H%M%S}.xlsx'
    if output.resolve() == args.template.resolve() or output.exists():
        parser.error('输出文件已存在或与模板相同，请选择新输出路径')
    reader = PDFReader(BASE/'.ocr-cache')
    customers = []
    for folder in folders:
        print(f'处理客户：{folder.name}', flush=True)
        main_fields, auth = collect_customer(folder, reader)
        customers.append((folder, main_fields, auth))
    audit = write_result(args.template, output, customers)
    output.with_suffix('.核对.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), 'utf-8')
    pending = [item for item in audit if item['status'] == '待核对']
    summary = [f'登记表生成成功：{output.resolve()}', f'客户数：{len(customers)}；待核对字段：{len(pending)}']
    if pending:
        summary.extend(['以下字段已留空并标黄，请核对原始材料后补填：', ''])
    summary.extend(
        f'【待核对】{item["客户文件夹"]}\n'
        f'  工作表：{item["工作表"]}；单元格：{item["单元格"]}；字段：{item["字段"]}\n'
        f'  原因：{item["note"]}'
        for item in pending
    )
    output.with_suffix('.待核对.txt').write_text('\n'.join(summary), 'utf-8')
    print('\n'.join(summary), flush=True)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
