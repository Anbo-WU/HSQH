"""规则与错误取值回归测试；不需要启动 OCR。"""
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock
from dataclasses import replace

import pymupdf
from openpyxl import Workbook, load_workbook

from 登记 import (beneficiary_names, capital_wan, credit_class, dates, expiry,
                financial_value, question18, suitability, continuation_headers,
                credit_score, collect_customer, write_result, read_kyc, read_authorizations,
                beneficiary_fields, beneficiary_validity, Field, BASE, SHEET1, SHEET2)
from ocr_support import Line, PDFReader


def line(text, x, y, width=.07):
    return Line(text, .99, x, y-.006, x+width, y+.006)


def beneficiary(name, y=.25, identity='11010519491231002X', validity='2021.08.11-2041.08.11'):
    result = [line('*股东/控制人/高管', .06, y, .14), line('*姓名', .21, y),
              line('*身份证件号码', .55, y, .12),
              line('*身份证件有限期', .55, y+.025, .12)]
    if name:
        result.append(line(name, .33, y))
    if identity:
        result.append(line(identity, .70, y, .2))
    if validity:
        result.append(line(validity, .70, y+.025, .25))
    return result


def authorization_sheet():
    book = Workbook()
    sheet = book.active
    sheet.title = '交易授权'
    # 故意留一空行；相同的手机号和邮箱也要逐人保留。
    for row, name in [(14, '张三'), (16, '李四')]:
        for col, value in {'B': name, 'H': 13800000000, 'E': '11010519491231002X',
                           'F': '2021.08.11-2041.08.11', 'I': 'shared@example.com'}.items():
            sheet[f'{col}{row}'] = value
    return sheet


class RegistrationTests(unittest.TestCase):
    def test_authorization_columns_share_rows_and_keep_duplicates(self):
        fields = read_authorizations(authorization_sheet(), 'KYC')
        self.assertEqual(fields['C'].value, '张三，李四')
        self.assertEqual(fields['E'].value, '13800000000，13800000000')
        self.assertEqual(fields['H'].value, 'shared@example.com，shared@example.com')
        for field in fields.values():
            self.assertEqual(field.status, '已提取')
            self.assertEqual(len(field.value.split('，')), 2)
            self.assertIn('授权人 2', field.note)
        self.assertIn('H16', fields['E'].evidence)

    def test_authorization_missing_middle_value_does_not_shift_next_person(self):
        sheet = authorization_sheet()
        sheet['H14'] = None
        fields = read_authorizations(sheet, 'KYC')
        self.assertIsNone(fields['E'].value)
        self.assertEqual(fields['E'].status, '待核对')
        self.assertIn('张三（第14行）', fields['E'].note)
        self.assertIn('手机号 1', fields['E'].note)
        self.assertEqual(fields['C'].value, '张三，李四')
        self.assertEqual(fields['H'].status, '已提取')

    def test_authorization_equal_counts_but_different_rows_are_rejected(self):
        sheet = authorization_sheet()
        sheet['H14'] = None
        sheet['H15'] = 13900000000
        fields = read_authorizations(sheet, 'KYC')
        for field in fields.values():
            self.assertEqual(field.status, '待核对')
            self.assertIsNone(field.value)
            self.assertIn('姓名为空或错误：15', field.note)

    def test_authorization_errors_and_numeric_id_precision(self):
        for column, value in [('E', 110105194912310000), ('I', '#VALUE!'), ('F', '=A1')]:
            with self.subTest(column=column):
                with TemporaryDirectory() as tmp:
                    folder = Path(tmp)
                    (folder/'0.OA').mkdir()
                    sheet = authorization_sheet()
                    sheet[f'{column}16'] = value
                    sheet.parent.save(folder/'0.OA/00.KYC.xlsx')
                    _, fields = read_kyc(folder)
                    target = {'E': 'F', 'I': 'H', 'F': 'G'}[column]
                    self.assertIsNone(fields[target].value)
                    self.assertIn(f'{column}16', fields[target].note)

    def test_qualification_and_suitability_copy(self):
        with TemporaryDirectory() as tmp:
            folder = Path(tmp)/'001、测试公司-C5'
            (folder/'0.OA').mkdir(parents=True)
            book = Workbook()
            book.active.title = '基本信息表'
            book.active['C8'] = '民营企业'
            book.save(folder/'0.OA/00.KYC.xlsx')
            fields, _ = collect_customer(folder, Mock())
            self.assertEqual(fields['L'].value, '民营企业')
            self.assertIn('C8', fields['L'].source)
            self.assertEqual(fields['AK'].value, fields['BA'].value)
            self.assertEqual(fields['AK'].value, 'C5')

    def test_beneficiary_ids_and_dates_follow_names_across_pages(self):
        pages = [beneficiary('张三')+beneficiary('', .55, '', ''),
                 beneficiary('李四', validity='2022年1月1日至长期')+beneficiary('张三', .55)]
        fields = beneficiary_fields(pages, '13.pdf')
        self.assertEqual(fields['R'].value, '张三、李四')
        self.assertEqual(fields['S'].value, '11010519491231002X、11010519491231002X')
        self.assertEqual(fields['T'].value, '2021.08.11-2041.08.11、2022.01.01-长期')

    def test_beneficiary_missing_id_does_not_take_next_person(self):
        fields = beneficiary_fields([beneficiary('张三', identity='')+beneficiary('李四', .55)], '13.pdf')
        self.assertIsNone(fields['S'].value)
        self.assertIn('张三', fields['S'].note)
        self.assertEqual(fields['R'].value, '张三、李四')
        self.assertEqual(fields['T'].status, '已提取')

    def test_beneficiary_conflicts_low_confidence_and_invalid_id(self):
        fields = beneficiary_fields([beneficiary('张三')+beneficiary('张三', .55, validity='2022.01.01-2042.01.01')], '13.pdf')
        self.assertEqual(fields['S'].status, '已提取')
        self.assertIsNone(fields['T'].value)
        for identity in ['110105194912310021', '11010519491231002', '11010519491331002X']:
            self.assertIsNone(beneficiary_fields([beneficiary('张三', identity=identity)], '13.pdf')['S'].value)
        entries = beneficiary('张三')
        entries = [replace(x, score=.5) if x.text == '张三' else x for x in entries]
        fields = beneficiary_fields([entries], '13.pdf')
        self.assertTrue(all(f.value is None for f in fields.values()))

    def test_beneficiary_wrapped_id_and_validity(self):
        entries = beneficiary('张三', identity='110105194912', validity='2021.08.11-')
        entries += [line('31002X', .70, .263, .2), line('2041.08.11', .70, .290, .2)]
        fields = beneficiary_fields([entries], '13.pdf')
        self.assertEqual(fields['S'].value, '11010519491231002X')
        self.assertEqual(fields['T'].value, '2021.08.11-2041.08.11')

    def test_beneficiary_blank_name_with_validity_is_not_ignored(self):
        with self.assertRaisesRegex(ValueError, '姓名未识别完整'):
            beneficiary_fields([beneficiary('张三')+beneficiary('', .55, '', '2021.01.01-2041.01.01')], '13.pdf')
        fields = beneficiary_fields([beneficiary('张三')+beneficiary('', .55, '', '')+
                                     [line('签署日期', .55, .9), line('2026.08.01', .70, .9)]], '13.pdf')
        self.assertEqual(fields['R'].value, '张三')

    def test_beneficiary_validity_validation(self):
        self.assertEqual(beneficiary_validity('长期'), '长期')
        for value in ['2021.02.30-2041.02.28', '2041.01.01-2021.01.01', '2041.01.01', '2021.01.01-2041.99.01']:
            with self.assertRaises(ValueError):
                beneficiary_validity(value)

    def test_output_formats_formulas_and_text_ids_for_multiple_customers(self):
        with TemporaryDirectory() as tmp:
            output = Path(tmp)/'result.xlsx'
            main = {'B': Field('测试公司'), 'AD': Field(1234567.891), 'AE': Field(123456.7),
                    'AF': Field(1111111.191), 'AG': Field(-1234.5), 'BC': Field(81.25),
                    'S': Field('11010519491231002X')}
            auth = read_authorizations(authorization_sheet(), 'KYC')
            write_result(BASE/'场外衍生品中心客户信息登记表-模板.xlsx', output,
                         [(Path('001、甲'), main, auth), (Path('002、乙'), main, auth)])
            book = load_workbook(output)
            for row in (3, 4):
                for col in ('AD', 'AE', 'AF', 'AG'):
                    self.assertEqual(book[SHEET1][f'{col}{row}'].number_format, '#,##0.00')
                    self.assertEqual(book[SHEET1][f'{col}{row}'].value, main[col].value)
                self.assertEqual(book[SHEET1][f'BC{row}'].number_format, '0.0')
                ratio = book[SHEET1][f'AH{row}']
                self.assertEqual(ratio.data_type, 'f')
                self.assertEqual(ratio.number_format, '0.00%')
                self.assertIn(f'AE{row}/AD{row}', ratio.value)
                self.assertNotIn('*100', ratio.value)
                self.assertIn('IFERROR', ratio.value)
                self.assertEqual(book[SHEET1][f'S{row}'].data_type, 's')
                self.assertEqual(book[SHEET2][f'F{row}'].data_type, 's')
                self.assertEqual(book[SHEET2][f'C{row}'].value, '张三，李四')
            book.close()

    def test_capital_units_and_chinese_numbers(self):
        for text, expected in [('壹仟万元整', 1000), ('壹亿贰仟万元整', 12000),
                               ('10000000元', 1000), ('1.5亿元', 15000),
                               ('人民币贰佰伍拾万元整', 250), ('拾万元整', 10),
                               ('壹仟贰佰叁拾肆万伍仟陆佰柒拾捌元整', 1234.5678)]:
            with self.subTest(text=text):
                self.assertEqual(capital_wan(text), expected)
        for text in ['1000', '壹仟万美元', '不清楚']:
            with self.assertRaises(ValueError):
                capital_wan(text)

    def test_expiry_and_invalid_dates(self):
        self.assertEqual(expiry('2021.08.11-2041.08.11'), '2041.08.11')
        self.assertEqual(expiry('2021-08-11至2041-08-11'), '2041.08.11')
        self.assertEqual(expiry('2021年8月11日至长期'), '长期')
        self.assertEqual(dates('226年8月31日 2026年2月30日'), [])
        with self.assertRaises(ValueError):
            expiry('看不清')
        with self.assertRaises(ValueError):
            expiry('2021.08.11-2041.99.11')

    def test_credit_boundaries(self):
        self.assertEqual(credit_class(80), '正常')
        self.assertEqual(credit_class(100), '正常')
        self.assertEqual(credit_class(77), '关注')
        self.assertEqual(credit_class(60.5), '关注')
        for value in [60, 59, 0, -1]:
            with self.assertRaises(ValueError):
                credit_class(value)

    def test_all_suitability_suffixes(self):
        for suffix in ['C4', 'C5', 'C4低买高', 'B类专业交易者']:
            self.assertEqual(suitability('396、某公司-'+suffix), suffix)
        self.assertEqual(suitability('396、某公司'), '专业交易者')

    def test_every_beneficiary_across_pages_blank_slots_and_duplicates(self):
        def person(name, y):
            return [line('*股东/控制人/高管', .06, y, .14),
                    line('*姓名', .21, y), line(name, .33, y),
                    line('*身份证件号码', .55, y, .12)]
        first = person('张三', .27)+person('李四', .45)
        second = person('王五', .27)+person('张三', .45)
        second += [line('*股东/控制人/高管', .06, .65, .14),
                   line('*姓名', .21, .65), line('*身份证件号码', .55, .65)]
        value, _ = beneficiary_names([first, second])
        self.assertEqual(value, '张三、李四、王五')

    def test_incomplete_beneficiary_is_reported(self):
        entries = [line('*股东/控制人/高管', .06, .27, .14), line('*姓名', .21, .27),
                   line('*身份证件号码', .55, .27), line('123456789012345678', .70, .27)]
        with self.assertRaises(ValueError):
            beneficiary_names([entries])

    def test_question_answer_only_at_stem_end(self):
        self.assertEqual(question18([line('18.贵单位认为自己能承受的最大投资损失是多少？B', .1, .88, .7)]), 'B')
        self.assertEqual(question18([line('18.最大投资损失是多少B', .1, .88, .7)]), 'B')
        self.assertEqual(question18([line('18.贵单位认为自己能承受的最大投资', .1, .85, .7),
                                     line('损失是多少？', .1, .88, .4), line('D', .8, .88)]), 'D')
        with self.assertRaises(ValueError):
            question18([line('18.最大投资损失是多少？', .1, .8, .5), line('B.√10%-30%', .1, .85)])

    def test_balance_columns_exclude_row_number_and_opening_balance(self):
        lines = [line('单位：元', .82, .1),
                 line('行次', .247, .14, .031), line('期末余额', .278, .14),
                 line('年初余额', .388, .14, .054),
                 line('行次', .654, .14, .032), line('期末余额', .710, .14, .054),
                 line('年初余额', .813, .14, .053),
                 line('资产合计', .111, .68, .053), line('30', .248, .68, .020),
                 line('8,146,053.66', .290, .68, .075),
                 line('2,563,901.29负债和所有者权益', .383, .68, .235),
                 line('53', .656, .68, .020), line('8,146,053.66', .713, .68, .075)]
        value, _ = financial_value(lines, '资产合计', '期末余额')
        self.assertEqual(value, 8146053.66)

    def test_profit_annual_column_negative_and_units(self):
        lines = [line('单位：万元', .80, .1), line('行次', .51, .18, .025),
                 line('本期金额', .60, .18, .06), line('本年累计金额', .75, .18, .10),
                 line('四、净利润（净亏损以负号填列）', .10, .89, .38),
                 line('31', .51, .89, .025), line('276.00', .61, .89, .07),
                 line('-3,020.15', .75, .89, .10)]
        value, _ = financial_value(lines, '净利润', '本年累计金额')
        self.assertEqual(value, -30201500)

    def test_scope_uses_kyc_even_when_license_is_missing(self):
        with TemporaryDirectory() as tmp:
            folder = Path(tmp)
            (folder/'0.OA').mkdir()
            for title in ('基本信息表', 'sheet1'):
                book = Workbook()
                book.active.title = title
                book.active['C14'] = 'KYC 经营范围；完整标点。'
                book.save(folder/'0.OA'/'00.KYC.xlsx')
                book.close()
                fields, _ = collect_customer(folder, Mock())
                self.assertEqual(fields['N'].value, 'KYC 经营范围；完整标点。')
                self.assertIn('C14', fields['N'].source)
                self.assertNotIn('营业执照', fields['N'].source)
            book = Workbook()
            book.active.title = 'sheet1'
            book.save(folder/'0.OA'/'00.KYC.xlsx')
            book.close()
            fields, _ = collect_customer(folder, Mock())
            self.assertIsNone(fields['N'].value)
            self.assertEqual(fields['N'].status, '待核对')

    def test_finance_alias_and_merged_neighbor_amount(self):
        entries = [line('单位：元', .8, .08), line('行次', .55, .15),
                   line('期末数', .69, .15), line('年初数', .82, .15),
                   line('0.00负债总计', .4, .7), line('47', .55, .7),
                   line('6,485,327.55', .69, .7), line('9,999,999.00', .82, .7)]
        self.assertEqual(financial_value(entries, '负债合计', '期末余额')[0], 6485327.55)
        entries[4] = line('流动负债合计', .4, .7)
        with self.assertRaises(ValueError):
            financial_value(entries, '负债合计', '期末余额')

    def test_missing_left_header_cannot_take_right_table_value(self):
        entries = [line('单位：元', .8, .08), line('期末余额', .7, .15),
                   line('年初余额', .82, .15), line('资产合计', .1, .7),
                   line('123.00', .3, .7), line('负债和所有者权益', .45, .7),
                   line('999.00', .7, .7)]
        with self.assertRaises(ValueError):
            financial_value(entries, '资产合计', '期末余额')

    def test_continuation_requires_aligned_consecutive_row_numbers(self):
        first = [line('单位：元', .8, .08), line('行次', .25, .15, .04),
                 line('期末余额', .4, .15), line('年初余额', .6, .15),
                 line('29', .25, .7, .04)]
        second = [line('资产总计', .1, .2), line('30', .25, .2, .04),
                  line('123.00', .4, .2), line('456.00', .6, .2)]
        inherited = continuation_headers(first, second)
        self.assertEqual(financial_value(second+inherited, '资产合计', '期末余额')[0], 123)
        for bad_row in (line('10', .25, .2, .04), line('30', .32, .2, .04)):
            self.assertEqual(continuation_headers(first, [second[0], bad_row]+second[2:]), [])
        self.assertEqual(continuation_headers(first, second+[line('利润表', .1, .1)]), [])

    def test_finance_unrecognized_unit_and_wrong_profit_period_are_rejected(self):
        entries = [line('净利润', .1, .7), line('本期金额', .4, .15),
                   line('123.00', .4, .7), line('单位：元', .8, .08)]
        with self.assertRaises(ValueError):
            financial_value(entries, '净利润', '本年累计金额')
        entries[1] = line('本年累计金额', .4, .15)
        for unit in ['单位：不详', '单位：美元']:
            with self.assertRaises(ValueError):
                financial_value(entries[:3]+[line(unit, .8, .08)], '净利润', '本年累计金额')

    def test_finance_scans_all_pages_and_marks_ambiguous_or_missing_values_yellow(self):
        balance = [line('单位：元', .8, .08), line('期末余额', .4, .15),
                   line('年初余额', .6, .15), line('资产总计', .1, .6),
                   line('1000.00', .4, .6), line('900.00', .6, .6),
                   line('负债合计', .1, .7), line('400.00', .4, .7), line('500.00', .6, .7)]
        profit = [line('单位：元', .8, .08), line('本年累计金额', .4, .15),
                  line('上年金额', .6, .15), line('四、净利润', .1, .6),
                  line('25.00', .4, .6), line('99.00', .6, .6)]
        with TemporaryDirectory() as tmp:
            folder = Path(tmp)
            (folder/'1.开户材料存档').mkdir()
            with pymupdf.open() as doc:
                for _ in range(4):
                    doc.new_page()
                doc.save(folder/'1.开户材料存档'/'16.财报表.pdf')
            reader = Mock()
            reader.oriented_page.side_effect = [([], 0), (profit, 270), ([], 0), (balance, 0)]
            fields, _ = collect_customer(folder, reader)
            self.assertEqual([fields[c].value for c in ('AD','AE','AF','AG')], [1000,400,600,25])
            self.assertIn('第4页', fields['AD'].source)
            reader.oriented_page.side_effect = [(balance, 0), (balance, 0), ([], 0), ([], 0)]
            fields, auth = collect_customer(folder, reader)
            output = folder/'result.xlsx'
            write_result(BASE/'场外衍生品中心客户信息登记表-模板.xlsx', output, [(folder, fields, auth)])
            book = load_workbook(output)
            for col in ('AD', 'AE', 'AF', 'AG'):
                self.assertIsNone(book[SHEET1][col+'3'].value)
                self.assertEqual(book[SHEET1][col+'3'].fill.fgColor.rgb, '00FFF2CC')
            book.close()

    def test_credit_local_direction_verification_and_real_disagreement(self):
        entries = [line('综合评分：', .08, .63), line('18', .22, .63),
                   line('综合评分81', .08, .74),
                   line('1.综合评分60以下为风险类', .64, .63)]
        reader = Mock()
        reader.page.return_value = [line('81', .22, .63)]
        field = credit_score(reader, Path('credit.pdf'), 1, entries)
        self.assertEqual(field.value, 81)
        self.assertIn('18 →', field.note)
        self.assertEqual(reader.page.call_count, 4)
        self.assertTrue(all(c.kwargs['use_cls'] is False for c in reader.page.call_args_list))
        reader.page.side_effect = [[line(v, .22, .63)] for v in ('81', '81', '82', '82')]
        field = credit_score(reader, Path('credit.pdf'), 1, entries)
        self.assertIsNone(field.value)
        self.assertEqual(field.status, '待核对')
        reader.page.side_effect = [[line(v, .22, .63)] for v in ('81', '18', '81', '81')]
        self.assertIsNone(credit_score(reader, Path('credit.pdf'), 1, entries).value)
        reader.page.side_effect = None
        reader.page.return_value = [line('81', .22, .63)]
        self.assertIsNone(credit_score(reader, Path('credit.pdf'), 1,
                                     entries+[line('综合评分：', .08, .85)]).value)

    def test_page_orientation_uses_text_instead_of_paper_size(self):
        reader = PDFReader(Path('unused'))
        vertical = [Line('资产总计', .99, .1, .1, .12, .3)]
        horizontal = [line('资产总计', .1, .3, .2), line('1000.00', .4, .3)]
        reader.page = Mock(side_effect=lambda *a, **kw: horizontal if kw.get('rotation') == 270 else vertical)
        self.assertEqual(reader.oriented_page(Path('finance.pdf'), 0, ('资产总计',)), (horizontal, 270))
        reader.page = Mock(return_value=horizontal)
        self.assertEqual(reader.oriented_page(Path('finance.pdf'), 0, ('资产总计',)), (horizontal, 0))
        self.assertEqual(reader.page.call_count, 1)

    def test_pdf_text_coordinates_follow_rotation_and_crop(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp)/'text.pdf'
            with pymupdf.open() as doc:
                page = doc.new_page(width=300, height=500)
                page.insert_text((40, 80), '81')
                doc.save(path)
            reader = PDFReader(Path(tmp)/'cache')
            original = reader.page(path, 0)[0]
            rotated = reader.page(path, 0, rotation=90)[0]
            self.assertAlmostEqual(rotated.x0, 1-original.y1, places=5)
            self.assertAlmostEqual(rotated.y0, original.x0, places=5)
            crop = (rotated.x0-.01, rotated.y0-.01, rotated.x1+.01, rotated.y1+.01)
            self.assertEqual(reader.page(path, 0, rotation=90, crop=crop)[0].text, '81')
            self.assertEqual(reader.page(path, 0, rotation=90, crop=(0, 0, .1, .1)), [])


if __name__ == '__main__':
    unittest.main()
