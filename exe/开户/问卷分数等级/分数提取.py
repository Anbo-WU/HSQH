#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从“问卷”文件夹的 Word/PDF 问卷中提取客户名称和 20 道题答案。"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

try:
    from openpyxl import load_workbook
except ImportError:
    print("错误：缺少 openpyxl。请先运行：python3 -m pip install openpyxl", file=sys.stderr)
    raise SystemExit(1)


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_DIR = BASE_DIR / "问卷"
DEFAULT_OUTPUT_FILE = BASE_DIR / "分数统计.xlsx"
QUESTION_COUNT = 20
MULTI_SELECT_QUESTIONS = {13, 17}
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{WORD_NAMESPACE}}}"

# 支持“1.”、“18．”、“1、”、“第1题式的第1.”以及“(1)”等常见写法。
QUESTION_RE = re.compile(
    r"^\s*(?:(?:第\s*)?(\d{1,2})\s*[.、:：]|[（(]\s*(\d{1,2})\s*[)）])"
)

# 答案必须在题目段落末尾，并位于冒号、问号或句号之后。
# 支持 AB、A B、A、B、A/B、A和B 等多选写法。
ANSWER_RE = re.compile(
    r"[:：?？。]\s*([A-Za-z](?:\s*(?:[,，、;/／+]|和)?\s*[A-Za-z])*)\s*$"
)
NAME_RE = re.compile(
    r"交易者名称\s*[:：]\s*(.*?)\s*证件号\s*[:：]", re.DOTALL
)

# macOS 上通过 PDFKit 读取文本型 PDF；如果某页没有文本层，则使用 Vision
# 做本地 OCR。源码由本脚本临时编译，不会在问卷目录生成额外文件。
PDF_HELPER_SOURCE = r'''
import Foundation
import PDFKit
import Vision
import AppKit

guard CommandLine.arguments.count == 2 else { exit(2) }
guard let document = PDFDocument(url: URL(fileURLWithPath: CommandLine.arguments[1])) else {
    fputs("无法打开 PDF\n", stderr)
    exit(3)
}

for pageNumber in 0..<document.pageCount {
    guard let page = document.page(at: pageNumber) else { continue }
    if let text = page.string,
       !text.trimmingCharacters(in: .whitespacesAndNewlines).isEmpty {
        print(text)
        continue
    }

    let bounds = page.bounds(for: .mediaBox)
    let scale: CGFloat = 2.0
    let width = Int(bounds.width * scale)
    let height = Int(bounds.height * scale)
    guard let bitmap = NSBitmapImageRep(
        bitmapDataPlanes: nil,
        pixelsWide: width,
        pixelsHigh: height,
        bitsPerSample: 8,
        samplesPerPixel: 4,
        hasAlpha: true,
        isPlanar: false,
        colorSpaceName: .deviceRGB,
        bytesPerRow: 0,
        bitsPerPixel: 0
    ), let context = NSGraphicsContext(bitmapImageRep: bitmap) else {
        fputs("无法渲染 PDF 第 \(pageNumber + 1) 页\n", stderr)
        exit(4)
    }

    NSGraphicsContext.saveGraphicsState()
    NSGraphicsContext.current = context
    context.cgContext.setFillColor(NSColor.white.cgColor)
    context.cgContext.fill(CGRect(x: 0, y: 0, width: width, height: height))
    context.cgContext.scaleBy(x: scale, y: scale)
    page.draw(with: .mediaBox, to: context.cgContext)
    NSGraphicsContext.restoreGraphicsState()

    guard let image = bitmap.cgImage else {
        fputs("无法生成 PDF 第 \(pageNumber + 1) 页图像\n", stderr)
        exit(5)
    }

    let request = VNRecognizeTextRequest()
    request.recognitionLevel = .accurate
    request.recognitionLanguages = ["zh-Hans", "en-US"]
    request.usesLanguageCorrection = false
    let handler = VNImageRequestHandler(cgImage: image, options: [:])
    do {
        try handler.perform([request])
    } catch {
        fputs("PDF 第 \(pageNumber + 1) 页 OCR 失败：\(error)\n", stderr)
        exit(6)
    }

    let observations = (request.results ?? []).sorted {
        if abs($0.boundingBox.midY - $1.boundingBox.midY) > 0.01 {
            return $0.boundingBox.midY > $1.boundingBox.midY
        }
        return $0.boundingBox.minX < $1.boundingBox.minX
    }
    for observation in observations {
        if let candidate = observation.topCandidates(1).first {
            print(candidate.string)
        }
    }
}
'''


class ExtractionError(Exception):
    """问卷内容不符合预期格式。"""


@dataclass(frozen=True)
class Questionnaire:
    source: Path
    customer_name: str
    answers: tuple[str, ...]


def normalize_text(text: str) -> str:
    """统一全角字符和不间断空格，同时保留段落边界。"""
    text = unicodedata.normalize("NFKC", text).replace("\u00a0", " ")
    return re.sub(r"[ \t]+", " ", text).strip()


def docx_paragraphs(path: Path) -> list[str]:
    """只用标准库读取 docx，避免额外依赖 python-docx。"""
    try:
        with ZipFile(path) as archive:
            xml_data = archive.read("word/document.xml")
    except (BadZipFile, KeyError, OSError) as exc:
        raise ExtractionError(f"无法读取 DOCX 文件：{exc}") from exc

    try:
        root = ET.fromstring(xml_data)
    except ET.ParseError as exc:
        raise ExtractionError(f"DOCX XML 已损坏：{exc}") from exc

    paragraphs: list[str] = []
    for paragraph in root.iter(W + "p"):
        parts: list[str] = []
        for node in paragraph.iter():
            if node.tag == W + "t" and node.text:
                parts.append(node.text)
            elif node.tag == W + "tab":
                parts.append("\t")
            elif node.tag in {W + "br", W + "cr"}:
                parts.append("\n")
        text = normalize_text("".join(parts))
        if text:
            paragraphs.append(text)
    return paragraphs


def decode_command_output(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "big5"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def legacy_doc_paragraphs(path: Path) -> list[str]:
    """读取旧版 .doc；macOS 优先使用系统自带的 textutil。"""
    commands: list[list[str]] = []
    if shutil.which("textutil"):
        commands.append(["textutil", "-convert", "txt", "-stdout", str(path)])
    if shutil.which("antiword"):
        commands.append(["antiword", str(path)])
    if shutil.which("catdoc"):
        commands.append(["catdoc", str(path)])

    if not commands:
        raise ExtractionError(
            "无法读取旧版 .doc：系统中没有 textutil、antiword 或 catdoc。"
            "请先用 Word/WPS 将文件另存为 .docx。"
        )

    failures: list[str] = []
    for command in commands:
        try:
            result = subprocess.run(
                command,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                check=False,
                timeout=60,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            failures.append(str(exc))
            continue

        if result.returncode == 0 and result.stdout:
            text = decode_command_output(result.stdout).replace("\r\n", "\n").replace("\r", "\n")
            return [normalize_text(line) for line in text.split("\n") if normalize_text(line)]
        failures.append(decode_command_output(result.stderr).strip() or f"退出码 {result.returncode}")

    raise ExtractionError("无法转换旧版 .doc：" + "；".join(failures))


def compile_pdf_helper() -> Path:
    """在系统临时目录编译并缓存 macOS PDF/OCR 小工具。"""
    if sys.platform != "darwin":
        raise ExtractionError(
            "扫描版 PDF 的自动识别目前需要 macOS；"
            "请将 PDF 转为可搜索 PDF 或 DOCX 后再运行。"
        )

    swiftc = shutil.which("swiftc")
    if not swiftc:
        raise ExtractionError(
            "系统中没有 swiftc，无法启用 PDF/OCR。"
            "请安装 Apple Command Line Tools，或将 PDF 转为 DOCX。"
        )

    source_hash = hashlib.sha256(PDF_HELPER_SOURCE.encode("utf-8")).hexdigest()[:16]
    cache_dir = Path(tempfile.gettempdir()) / "questionnaire_pdf_reader"
    executable = cache_dir / f"pdf_reader_{source_hash}"
    if executable.is_file() and os.access(executable, os.X_OK):
        return executable

    cache_dir.mkdir(parents=True, exist_ok=True)
    module_cache = cache_dir / "module_cache"
    module_cache.mkdir(exist_ok=True)
    source_file: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            prefix="pdf_reader_",
            suffix=".swift",
            dir=cache_dir,
            delete=False,
        ) as stream:
            stream.write(PDF_HELPER_SOURCE)
            source_file = Path(stream.name)

        environment = os.environ.copy()
        environment["SWIFT_MODULECACHE_PATH"] = str(module_cache)
        environment["CLANG_MODULE_CACHE_PATH"] = str(module_cache)
        result = subprocess.run(
            [swiftc, str(source_file), "-o", str(executable)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=False,
            timeout=120,
            env=environment,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise ExtractionError(f"编译 PDF 识别组件失败：{exc}") from exc
    finally:
        if source_file is not None:
            source_file.unlink(missing_ok=True)

    if result.returncode != 0 or not executable.is_file():
        detail = decode_command_output(result.stderr).strip()
        raise ExtractionError(f"编译 PDF 识别组件失败：{detail or '未知错误'}")
    return executable


def pdf_paragraphs(path: Path) -> list[str]:
    """读取 PDF；没有文本层的页面会自动使用中文 OCR。"""
    helper = compile_pdf_helper()
    try:
        result = subprocess.run(
            [str(helper), str(path.resolve())],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=False,
            timeout=180,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise ExtractionError(f"PDF 读取/OCR 失败：{exc}") from exc

    if result.returncode != 0:
        detail = decode_command_output(result.stderr).strip()
        raise ExtractionError(f"PDF 读取/OCR 失败：{detail or f'退出码 {result.returncode}'}")

    text = decode_command_output(result.stdout).replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [normalize_text(line) for line in text.split("\n") if normalize_text(line)]
    if not paragraphs:
        raise ExtractionError("PDF 中没有识别到文字")
    return paragraphs


def read_paragraphs(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return docx_paragraphs(path)
    if suffix == ".doc":
        return legacy_doc_paragraphs(path)
    if suffix == ".pdf":
        return pdf_paragraphs(path)
    raise ExtractionError(f"不支持的文件格式：{suffix}")


def extract_customer_name(paragraphs: Iterable[str]) -> str:
    full_text = "\n".join(paragraphs)
    match = NAME_RE.search(full_text)
    if not match:
        raise ExtractionError("未找到“交易者名称：……证件号：”字段")

    name = re.sub(r"\s+", " ", match.group(1)).strip()
    if not name:
        raise ExtractionError("交易者名称为空")
    return name


def question_number(paragraph: str) -> int | None:
    match = QUESTION_RE.match(paragraph)
    if not match:
        return None
    return int(match.group(1) or match.group(2))


def extract_answer(paragraph: str, number: int) -> str:
    match = ANSWER_RE.search(paragraph)
    if not match:
        raise ExtractionError(f"第 {number} 题题目段落末尾未找到答案：{paragraph}")

    letters = re.findall(r"[A-Za-z]", match.group(1).upper())
    if not letters:
        raise ExtractionError(f"第 {number} 题答案为空")
    if len(set(letters)) != len(letters):
        raise ExtractionError(f"第 {number} 题答案含重复选项：{''.join(letters)}")
    if number not in MULTI_SELECT_QUESTIONS and len(letters) != 1:
        raise ExtractionError(
            f"第 {number} 题是单选题，但检测到多个答案：{''.join(letters)}"
        )
    return "".join(letters)


def join_wrapped_questions(paragraphs: list[str]) -> list[str]:
    """合并 PDF/OCR 中被自动换行的题目，直到读到段尾答案。"""
    joined: list[str] = []
    index = 0
    while index < len(paragraphs):
        paragraph = paragraphs[index]
        number = question_number(paragraph)
        if number is None or not 1 <= number <= QUESTION_COUNT:
            joined.append(paragraph)
            index += 1
            continue

        parts = [paragraph]
        next_index = index + 1
        while not ANSWER_RE.search("".join(parts)) and next_index < len(paragraphs):
            following = paragraphs[next_index]
            if question_number(following) is not None:
                break
            # OCR 偶尔把页码识别成单独一行，题目合并时忽略它。
            if not re.fullmatch(r"\d{1,5}", following):
                parts.append(following)
            next_index += 1
        joined.append("".join(parts))
        index = next_index
    return joined


def extract_questionnaire(path: Path) -> Questionnaire:
    paragraphs = read_paragraphs(path)
    if not paragraphs:
        raise ExtractionError("文档中没有可读取的文字")

    customer_name = extract_customer_name(paragraphs)
    paragraphs = join_wrapped_questions(paragraphs)
    answers: dict[int, str] = {}

    for paragraph in paragraphs:
        number = question_number(paragraph)
        if number is None or not 1 <= number <= QUESTION_COUNT:
            continue
        if number in answers:
            raise ExtractionError(f"检测到重复的第 {number} 题题目段落")
        answers[number] = extract_answer(paragraph, number)

    missing = [number for number in range(1, QUESTION_COUNT + 1) if number not in answers]
    if missing:
        missing_text = "、".join(map(str, missing))
        raise ExtractionError(f"缺少第 {missing_text} 题或这些题未被正确识别")

    ordered_answers = tuple(answers[number] for number in range(1, QUESTION_COUNT + 1))
    return Questionnaire(path, customer_name, ordered_answers)


def natural_sort_key(path: Path) -> list[object]:
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", path.name)]


def find_questionnaire_files(input_dir: Path) -> list[Path]:
    if not input_dir.is_dir():
        raise ExtractionError(f"找不到问卷文件夹：{input_dir}")

    files = [
        path
        for path in input_dir.iterdir()
        if path.is_file()
        and path.suffix.lower() in {".doc", ".docx", ".pdf"}
        and not path.name.startswith((".~", "~$", "."))
    ]
    return sorted(files, key=natural_sort_key)


def validate_headers(worksheet) -> None:
    expected = ["客户名称", *range(1, QUESTION_COUNT + 1)]
    actual = [worksheet.cell(1, column).value for column in range(1, 22)]
    if actual != expected:
        raise ExtractionError(
            "分数统计.xlsx 的 A1:U1 表头与预期不一致；应为“客户名称、1、2、……、20”"
        )


def write_workbook(output_file: Path, records: list[Questionnaire]) -> None:
    if not output_file.is_file():
        raise ExtractionError(f"找不到 Excel 文件：{output_file}")

    try:
        workbook = load_workbook(output_file)
    except Exception as exc:
        raise ExtractionError(f"无法打开 Excel 文件：{exc}") from exc

    worksheet = workbook.active
    validate_headers(worksheet)

    # 仅替换 A:U 的旧数据，V、W 列的内容、公式和格式均不改动。
    last_row = max(worksheet.max_row, len(records) + 1)
    for row in worksheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=21):
        for cell in row:
            cell.value = None

    for row_number, record in enumerate(records, start=2):
        worksheet.cell(row_number, 1, record.customer_name)
        for column, answer in enumerate(record.answers, start=2):
            worksheet.cell(row_number, column, answer)

    output_file = output_file.resolve()
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_file.stem}_",
            suffix=output_file.suffix,
            dir=output_file.parent,
            delete=False,
        ) as temp_file:
            temp_path = Path(temp_file.name)
        workbook.save(temp_path)
        os.replace(temp_path, output_file)
    except Exception as exc:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)
        raise ExtractionError(
            f"保存 Excel 失败（请确认文件未被 Word/WPS/Excel 占用）：{exc}"
        ) from exc


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="提取 Word/PDF 问卷中的客户名称及 20 道题答案，写入分数统计.xlsx。"
    )
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR, help="问卷文件夹路径")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_FILE, help="Excel 文件路径")
    parser.add_argument("--preview", action="store_true", help="只检查并显示结果，不修改 Excel")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        files = find_questionnaire_files(args.input_dir)
        if not files:
            raise ExtractionError(f"“{args.input_dir}”中没有找到 .doc、.docx 或 .pdf 文件")

        records: list[Questionnaire] = []
        errors: list[str] = []
        for path in files:
            try:
                records.append(extract_questionnaire(path))
            except ExtractionError as exc:
                errors.append(f"{path.name}：{exc}")

        if errors:
            print("提取失败，Excel 未作任何修改：", file=sys.stderr)
            for error in errors:
                print(f"  - {error}", file=sys.stderr)
            return 1

        print(f"已成功检查 {len(records)} 份问卷：")
        for record in records:
            print(f"  {record.source.name} -> {record.customer_name} | {' '.join(record.answers)}")

        if args.preview:
            print("当前为预览模式，未修改 Excel。")
        else:
            write_workbook(args.output, records)
            print(f"已写入：{args.output.resolve()}")
        return 0
    except ExtractionError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
