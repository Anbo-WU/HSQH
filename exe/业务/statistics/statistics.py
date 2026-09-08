#!/usr/bin/env python3
"""扫描确认书或结算单 PDF，并把结果写入各自目录中的统计表。

默认行为：
1. 选择 A 时扫描“确认书”目录，选择 B 时扫描“结算单”目录；
2. 自动定位表头行和现有第一条参考数据，从第二条数据行开始写入；
3. 不覆盖源统计表，生成“保期数据统计表_已填充.xlsx”；
4. 如果目标单元格已有内容，除非显式使用 --force，否则停止并提示。

常用命令：
    python3 statistics.py A --preview
    python3 statistics.py A
    python3 statistics.py B --preview
    python3 statistics.py B --output 自定义结果.xlsx
    python3 statistics.py --self-test

Windows 版依赖 openpyxl、pypdf、PyMuPDF、RapidOCR 和 ONNX Runtime。
PDF 有文字层时优先使用 pypdf；纯图片扫描件自动使用本地 OCR。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
import unicodedata
from copy import copy
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook


TWO_PLACES = Decimal("0.01")
CONFIRMATION_COLUMNS = {
    "U": "对冲方式",
    "V": "入场价格",
    "W": "名义数量",
    "X": "名义本金",
    "Y": "权利金",
    "Z": "场外期权型",
    "AA": "对冲比例",
    "AB": "对冲标的合约",
    "AC": "入场时间",
    "AD": "交易确认书生效日",
    "AE": "交易确认书到期日",
    "BB": "执行价格",
    "BN": "交易书结算价格",
    "AN": "交易确认书编号",
}
SETTLEMENT_COLUMNS = {
    "AO": "结算单编号",
    "BB": "执行价格",
    "BN": "交易书结算价格",
    "BF": "是否产生赔付",
    "BH": "期权赔付金额",
    "BJ": "理赔价格",
}
# 保留原名称，兼容已有调用方和测试。
TARGET_COLUMNS = CONFIRMATION_COLUMNS


@dataclass(frozen=True)
class ConfirmationData:
    source_file: str
    transaction_id: str
    hedge_method: str
    entry_price: Decimal
    nominal_quantity: Decimal
    nominal_principal: Decimal
    premium_total: Decimal
    otc_option_type: str
    hedge_ratio: Decimal
    contract: str
    signing_date: date
    effective_date: date
    expiry_date: date
    is_egg: bool
    exercise_price: Decimal | None
    pricing_description: str
    pricing_warning: str | None = None

    def preview_dict(self) -> dict[str, str]:
        return {
            "文件": self.source_file,
            "交易编号": self.transaction_id,
            "U": self.hedge_method,
            "V": f"{self.entry_price:.2f}",
            "W": f"{self.nominal_quantity:.2f}",
            "X": f"{self.nominal_principal:.2f}",
            "Y": f"{self.premium_total:.2f}",
            "Z": self.otc_option_type,
            "AA": "100.00%",
            "AB": self.contract,
            "AC": self.signing_date.isoformat(),
            "AD": self.effective_date.isoformat(),
            "AE": self.expiry_date.isoformat(),
            "BB": f"{self.exercise_price:.2f}" if self.exercise_price is not None else "",
            "BN": self.pricing_description,
            "AN": self.transaction_id,
        }


@dataclass(frozen=True)
class SettlementData:
    source_file: str
    settlement_id: str
    has_payout: bool
    payout_amount: Decimal
    settlement_price: Decimal
    exercise_price: Decimal | None
    pricing_description: str
    pricing_warning: str | None = None

    def preview_dict(self) -> dict[str, str]:
        return {
            "文件": self.source_file,
            "AO": self.settlement_id,
            "BF": "是" if self.has_payout else "否",
            "BH": f"{self.payout_amount:.2f}",
            "BJ": f"{self.settlement_price:.2f}",
            "BB": f"{self.exercise_price:.2f}" if self.exercise_price is not None else "",
            "BN": self.pricing_description,
        }


@dataclass(frozen=True)
class FailedDocument:
    source_file: str
    error_message: str

    @property
    def note(self) -> str:
        return f"识别异常：{self.source_file}：{self.error_message}"


@dataclass(frozen=True)
class PricingTerms:
    option_direction: str
    description: str
    exercise_price: Decimal | None
    warning: str | None = None


ConfirmationRecord = ConfirmationData | FailedDocument
SettlementRecord = SettlementData | FailedDocument


class PDFTextExtractor:
    """优先提取文字层；必要时在一次运行中复用 Windows 本地 OCR。"""

    def __init__(self) -> None:
        self._ocr_engine: object | None = None

    def close(self) -> None:
        self._ocr_engine = None

    def __enter__(self) -> "PDFTextExtractor":
        return self

    def __exit__(self, *_: object) -> None:
        self.close()

    @staticmethod
    def _looks_usable(text: str) -> bool:
        if len(text.strip()) < 80:
            return False
        confirmation_markers = ("交易确认书", "生效日", "标的合约", "入场价格")
        settlement_markers = ("结算通知单", "结算日", "结算价", "结算金额")
        return (
            sum(marker in text for marker in confirmation_markers) >= 2
            or sum(marker in text for marker in settlement_markers) >= 2
        )

    def extract(self, pdf_path: Path) -> str:
        errors: list[str] = []

        try:
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(str(pdf_path))
            text = "\n\n---PAGE---\n\n".join(page.extract_text() or "" for page in reader.pages)
            if self._looks_usable(text):
                return text
            errors.append("pypdf 未得到可用文字")
        except Exception as exc:  # pypdf 是可选依赖
            errors.append(f"pypdf: {exc}")

        try:
            text = self._extract_with_ocr(pdf_path)
            if self._looks_usable(text):
                return text
            errors.append("RapidOCR 未得到可用文字")
        except Exception as exc:
            errors.append(f"RapidOCR: {exc}")

        details = "；".join(errors)
        raise RuntimeError(
            f"无法识别 PDF：{pdf_path.name}。{details}。"
            "请确认已安装 pypdf、PyMuPDF、RapidOCR 和 ONNX Runtime。"
        )

    def _ensure_ocr_engine(self) -> object:
        if self._ocr_engine is not None:
            return self._ocr_engine
        try:
            from rapidocr import RapidOCR  # type: ignore
        except ImportError as exc:
            raise RuntimeError("缺少 rapidocr") from exc
        self._ocr_engine = RapidOCR()
        return self._ocr_engine

    @staticmethod
    def _ordered_ocr_text(result: object) -> str:
        import numpy as np

        texts = getattr(result, "txts", None)
        boxes = getattr(result, "boxes", None)
        if texts is None or boxes is None:
            return ""
        lines: list[tuple[float, float, str]] = []
        for text, box in zip(texts, boxes):
            points = np.asarray(box, dtype=float)
            lines.append(
                (
                    float(points[:, 1].min()),
                    float(points[:, 0].min()),
                    str(text),
                )
            )
        lines.sort(key=lambda item: (item[0], item[1]))
        return "\n".join(text for _, _, text in lines)

    def _extract_with_ocr(self, pdf_path: Path) -> str:
        try:
            import numpy as np
            import pymupdf
        except ImportError as exc:
            raise RuntimeError("缺少 PyMuPDF 或 numpy") from exc

        engine = self._ensure_ocr_engine()
        try:
            document = pymupdf.open(pdf_path)
        except Exception as exc:
            raise RuntimeError("无法打开 PDF") from exc
        pages: list[str] = []
        try:
            if document.page_count < 1:
                raise RuntimeError("PDF 没有页面")
            print(
                f"{pdf_path.name} 没有可用文字层，改用 Windows OCR……",
                file=sys.stderr,
            )
            for page_index in range(document.page_count):
                pixmap = document[page_index].get_pixmap(
                    matrix=pymupdf.Matrix(2.5, 2.5),
                    colorspace=pymupdf.csRGB,
                    alpha=False,
                )
                image = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
                    pixmap.height,
                    pixmap.width,
                    pixmap.n,
                )
                try:
                    result = engine(image)
                except Exception as exc:
                    raise RuntimeError(f"第 {page_index + 1} 页 OCR 失败") from exc
                pages.append(self._ordered_ocr_text(result))
                print(
                    f"OCR 进度：{page_index + 1}/{document.page_count}",
                    file=sys.stderr,
                )
        finally:
            document.close()
        return "\n\n---PAGE---\n\n".join(pages)


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\u00a0", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return text


def require_match(pattern: str, text: str, field_name: str, flags: int = 0) -> str:
    match = re.search(pattern, text, flags)
    if not match:
        raise ValueError(f"未找到字段：{field_name}")
    return match.group(1).strip()


def parse_decimal(raw: str, field_name: str) -> Decimal:
    try:
        return Decimal(raw.replace(",", "").strip()).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)
    except InvalidOperation as exc:
        raise ValueError(f"{field_name}不是有效数值：{raw!r}") from exc


def parse_date(raw: str, field_name: str) -> date:
    numbers = re.findall(r"\d+", raw)
    if len(numbers) < 3:
        raise ValueError(f"{field_name}不是有效日期：{raw!r}")
    try:
        return date(int(numbers[0]), int(numbers[1]), int(numbers[2]))
    except ValueError as exc:
        raise ValueError(f"{field_name}不是有效日期：{raw!r}") from exc


def clean_label_value(raw: str) -> str:
    return re.sub(r"[\s【】\[\]]+", "", raw).strip(" :：")


def clean_transaction_id(raw: str) -> str:
    """清理交易编号；PDF 把后续标签粘到同一行时，在首个汉字处结束。"""
    before_chinese = re.split(r"[\u3400-\u4dbf\u4e00-\u9fff]", raw, maxsplit=1)[0]
    transaction_id = re.sub(r"\s+", "", before_chinese).strip(" :：")
    if not transaction_id:
        raise ValueError("交易编号不是有效内容")
    return transaction_id


def extract_option_type(text: str, section_mark: str) -> tuple[str, str]:
    option_type = clean_label_value(
        require_match(
            r"3" + section_mark + r"1\s*[、,]?\s*期权类型\s*[:：]\s*"
            r"([^\n]*?)(?=\s*3" + section_mark + r"2\s*[、,]?|\n|$)",
            text,
            "期权类型",
        )
    )
    direction_match = re.search(r"看涨|看跌", option_type)
    if not direction_match:
        raise ValueError(f"期权类型中未找到看涨或看跌：{option_type!r}")
    return option_type, direction_match.group(0)


def clean_pricing_description(raw: str) -> str:
    """去掉 PDF 断行造成的拆字，并保留便于阅读的 Max/Min 与单位空格。"""
    description = re.sub(r"\s+", "", raw).strip(" :：;；")
    description = re.sub(r"(?<![A-Za-z])(Max|Min)(?=【)", r" \1", description)
    description = re.sub(r"(?<=\d)(?=元/吨)", " ", description)
    return description


def parse_pricing_terms(text: str, option_direction: str) -> PricingTerms:
    """按 3.1 的看涨/看跌方向，从 5.4 选择对应结算价格描述及括号数值。"""
    section_match = re.search(
        r"5\s*[.．]\s*4\s*[、,]?\s*结算价格(?:计算方式)?\s*[:：]\s*"
        r"([\s\S]*?)(?=\s*5\s*[.．]\s*5\s*[、,]?|$)",
        text,
    )
    if not section_match:
        return PricingTerms(
            option_direction=option_direction,
            description="",
            exercise_price=None,
            warning="未找到 5.4 结算价格描述",
        )

    section = section_match.group(1).strip()
    branch_pattern = re.compile(
        r"(?:[12]\s*[、.)．]?\s*)?(看涨|看跌)(?:期权)?\s*[:：]"
    )
    branches = list(branch_pattern.finditer(section))
    selected = ""
    warnings: list[str] = []
    allow_numeric_extraction = True

    matching_index = next(
        (index for index, branch in enumerate(branches) if branch.group(1) == option_direction),
        None,
    )
    if matching_index is not None:
        branch = branches[matching_index]
        end = (
            branches[matching_index + 1].start()
            if matching_index + 1 < len(branches)
            else len(section)
        )
        selected = section[branch.end():end]
    elif branches:
        selected = section
        allow_numeric_extraction = False
        found = "、".join(branch.group(1) for branch in branches)
        warnings.append(
            f"5.4 中没有与 3.1“{option_direction}”匹配的分支（现有：{found}）"
        )
    else:
        pricing_start = section.find("采价期间")
        selected = section[pricing_start:] if pricing_start >= 0 else section
        warnings.append("5.4 不是标准的看涨/看跌分支结构，已保留可识别描述")

    description = clean_pricing_description(selected)
    if not description:
        description = clean_pricing_description(section)
        warnings.append("5.4 选中分支没有可写入的描述")

    exercise_price: Decimal | None = None
    if allow_numeric_extraction:
        compact = re.sub(r"\s+", "", selected)
        number_match = re.search(
            r"【[^】]*?[,，]([-+]?\d[\d,]*(?:\.\d+)?)元/吨】",
            compact,
        )
        if number_match:
            exercise_price = parse_decimal(number_match.group(1), "5.4 执行价格")
        else:
            warnings.append("5.4 描述中未找到“【...,数值 元/吨】”，BB 列将留空")

    return PricingTerms(
        option_direction=option_direction,
        description=description,
        exercise_price=exercise_price,
        warning="；".join(dict.fromkeys(warnings)) or None,
    )


def parse_confirmation(text: str, source_file: str) -> ConfirmationData:
    text = normalize_text(text)
    section_mark = r"\s*[.．]\s*"
    date_pattern = r"(\d{4}\s*(?:年|[-/.])\s*\d{1,2}\s*(?:月|[-/.])\s*\d{1,2}\s*日?)"

    transaction_id = clean_transaction_id(
        require_match(r"交易编号\s*[:：]\s*([^\n]+)", text, "交易编号")
    )
    signing_date = parse_date(
        require_match(r"签订时间\s*[:：]\s*" + date_pattern, text, "签订时间"),
        "签订时间",
    )
    effective_date = parse_date(
        require_match(r"2" + section_mark + r"1\s*[、,]?\s*生效日\s*[:：]\s*" + date_pattern,
                      text, "生效日"),
        "生效日",
    )
    expiry_date = parse_date(
        require_match(r"2" + section_mark + r"2\s*[、,]?\s*到期日\s*[:：]\s*" + date_pattern,
                      text, "到期日"),
        "到期日",
    )
    option_type, option_direction = extract_option_type(text, section_mark)
    pricing_terms = parse_pricing_terms(text, option_direction)
    option_structure = clean_label_value(
        require_match(
            r"3" + section_mark + r"2\s*[、,]?\s*期权结构\s*[:：]\s*"
            r"([^\n]*?)(?=\s*4\s*[、,]?|\n|$)",
            text,
            "期权结构",
        )
    )
    contract = require_match(
        r"4" + section_mark + r"1\s*[、,]?\s*标的合约\s*[:：]\s*([A-Za-z]+\s*[-]?\s*\d{3,4})",
        text,
        "标的合约",
    )
    contract = re.sub(r"\s+", "", contract).upper()
    quantity = parse_decimal(
        require_match(
            r"4" + section_mark + r"1\s*[、,]?[\s\S]{0,160}?数量\s*[:：]\s*([0-9][0-9,]*(?:\.\d+)?)",
            text,
            "数量",
        ),
        "数量",
    )
    entry_price = parse_decimal(
        require_match(
            r"5" + section_mark + r"1\s*[、,]?\s*入场价格\s*[:：]\s*([0-9][0-9,]*(?:\.\d+)?)",
            text,
            "入场价格",
        ),
        "入场价格",
    )
    premium_total = parse_decimal(
        require_match(
            r"5" + section_mark + r"5\s*[、,]?\s*期权费\s*[:：][\s\S]{0,260}?"
            r"合计\s*[:：]\s*([0-9][0-9,]*(?:\.\d+)?)",
            text,
            "期权费合计",
        ),
        "期权费合计",
    )

    is_egg = "鸡蛋" in text or bool(re.match(r"^JD\d", contract, re.IGNORECASE))
    if is_egg:
        entry_price = (entry_price * Decimal("2")).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)
    principal = (entry_price * quantity).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)
    otc_option_type = f"{option_structure}{option_type}"

    if quantity <= 0 or entry_price <= 0 or premium_total < 0:
        raise ValueError("数量、入场价格必须为正数，权利金不能为负数")
    if not (signing_date <= effective_date <= expiry_date):
        raise ValueError(
            f"日期顺序异常：签订日 {signing_date}、生效日 {effective_date}、到期日 {expiry_date}"
        )

    return ConfirmationData(
        source_file=source_file,
        transaction_id=transaction_id,
        hedge_method="期货",
        entry_price=entry_price,
        nominal_quantity=quantity,
        nominal_principal=principal,
        premium_total=premium_total,
        otc_option_type=otc_option_type,
        hedge_ratio=Decimal("1.00"),
        contract=contract,
        signing_date=signing_date,
        effective_date=effective_date,
        expiry_date=expiry_date,
        is_egg=is_egg,
        exercise_price=pricing_terms.exercise_price,
        pricing_description=pricing_terms.description,
        pricing_warning=pricing_terms.warning,
    )


def parse_settlement(text: str, source_file: str) -> SettlementData:
    """提取结算单编号、赔付状态/金额和结算价。"""
    text = normalize_text(text)
    section_mark = r"\s*[.．]\s*"
    number_pattern = r"([-+−]?\s*[0-9][0-9,]*(?:\.\d+)?)"

    _, option_direction = extract_option_type(text, section_mark)
    pricing_terms = parse_pricing_terms(text, option_direction)

    # 只匹配独占行首的“编号”，避免误取正文中的“交易确认书编号”。
    settlement_id = clean_transaction_id(
        require_match(r"(?m)^\s*编号\s*[:：]\s*([^\n]+)", text, "结算单编号")
    )
    settlement_price = parse_decimal(
        require_match(
            r"6" + section_mark + r"2\s*[、,]?\s*结算价\s*[:：]\s*" + number_pattern,
            text,
            "结算价",
        ).replace("−", "-").replace(" ", ""),
        "结算价",
    )
    raw_amount = parse_decimal(
        require_match(
            r"6" + section_mark + r"4\s*[、,]?\s*结算金额\s*[:：]\s*" + number_pattern,
            text,
            "结算金额",
        ).replace("−", "-").replace(" ", ""),
        "结算金额",
    )
    has_payout = raw_amount > 0
    payout_amount = raw_amount if has_payout else Decimal("0.00")

    return SettlementData(
        source_file=source_file,
        settlement_id=settlement_id,
        has_payout=has_payout,
        payout_amount=payout_amount,
        settlement_price=settlement_price,
        exercise_price=pricing_terms.exercise_price,
        pricing_description=pricing_terms.description,
        pricing_warning=pricing_terms.warning,
    )


def discover_pdfs(folder: Path) -> list[Path]:
    pdfs = [
        path
        for path in folder.iterdir()
        if path.is_file()
        and path.suffix.lower() == ".pdf"
        and "模板确认书" not in path.stem
    ]
    return sorted(pdfs, key=lambda path: path.name)


def extract_all(pdf_paths: Iterable[Path]) -> list[ConfirmationRecord]:
    results: list[ConfirmationRecord] = []
    seen_content: dict[str, str] = {}
    with PDFTextExtractor() as extractor:
        for pdf_path in pdf_paths:
            try:
                text = extractor.extract(pdf_path)
                fingerprint_source = re.sub(r"\s+", "", normalize_text(text))
                fingerprint = hashlib.sha256(fingerprint_source.encode("utf-8")).hexdigest()
                if fingerprint in seen_content:
                    print(
                        f"跳过重复 PDF：{pdf_path.name}（内容与 {seen_content[fingerprint]} 相同）",
                        file=sys.stderr,
                    )
                    continue
                seen_content[fingerprint] = pdf_path.name
                record = parse_confirmation(text, pdf_path.name)
                if record.pricing_warning:
                    print(
                        f"5.4 识别异常：{pdf_path.name}：{record.pricing_warning}",
                        file=sys.stderr,
                    )
                results.append(record)
            except Exception as exc:
                failure = FailedDocument(pdf_path.name, str(exc))
                results.append(failure)
                print(
                    f"确认书识别异常，已在统计表保留一行：{failure.note}",
                    file=sys.stderr,
                )
    return results


def extract_all_settlements(pdf_paths: Iterable[Path]) -> list[SettlementRecord]:
    results: list[SettlementRecord] = []
    seen_content: dict[str, str] = {}
    with PDFTextExtractor() as extractor:
        for pdf_path in pdf_paths:
            try:
                text = extractor.extract(pdf_path)
                fingerprint_source = re.sub(r"\s+", "", normalize_text(text))
                fingerprint = hashlib.sha256(fingerprint_source.encode("utf-8")).hexdigest()
                if fingerprint in seen_content:
                    print(
                        f"跳过重复 PDF：{pdf_path.name}（内容与 {seen_content[fingerprint]} 相同）",
                        file=sys.stderr,
                    )
                    continue
                seen_content[fingerprint] = pdf_path.name
                record = parse_settlement(text, pdf_path.name)
                if record.pricing_warning:
                    print(
                        f"5.4 识别异常：{pdf_path.name}：{record.pricing_warning}",
                        file=sys.stderr,
                    )
                results.append(record)
            except Exception as exc:
                failure = FailedDocument(pdf_path.name, str(exc))
                results.append(failure)
                print(
                    f"结算单识别异常，已在统计表保留一行：{failure.note}",
                    file=sys.stderr,
                )
    return results


def _cell_text(value: object) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def locate_rows_for_columns(
    ws: object, target_columns: dict[str, str], anchor_column: str
) -> tuple[int, int, int]:
    header_row = 0
    for row in range(1, min(getattr(ws, "max_row"), 30) + 1):
        if target_columns[anchor_column] in _cell_text(ws[f"{anchor_column}{row}"].value):
            header_row = row
            break
    if not header_row:
        raise ValueError(
            f"未在 {anchor_column} 列找到“{target_columns[anchor_column]}”表头，无法确认写入位置"
        )

    for column, keyword in target_columns.items():
        actual = _cell_text(ws[f"{column}{header_row}"].value)
        if keyword not in actual:
            raise ValueError(
                f"表头校验失败：{column}{header_row} 应包含“{keyword}”，实际为“{actual}”"
            )
    reference_row = header_row + 1
    start_row = header_row + 2
    return header_row, reference_row, start_row


def locate_rows(ws: object) -> tuple[int, int, int]:
    return locate_rows_for_columns(ws, CONFIRMATION_COLUMNS, "U")


def copy_reference_format(
    ws: object,
    reference_row: int,
    target_row: int,
    target_columns: Iterable[str] = CONFIRMATION_COLUMNS,
) -> None:
    for column in target_columns:
        source = ws[f"{column}{reference_row}"]
        target = ws[f"{column}{target_row}"]
        if source.has_style:
            target._style = copy(source._style)
    if ws.row_dimensions[reference_row].height is not None:
        ws.row_dimensions[target_row].height = ws.row_dimensions[reference_row].height


def write_workbook(
    workbook_path: Path,
    output_path: Path,
    records: Sequence[ConfirmationRecord],
    start_row_override: int | None = None,
    force: bool = False,
) -> tuple[str, int, int]:
    if workbook_path.suffix.lower() != ".xlsx":
        raise ValueError("当前统计表必须是 .xlsx；旧版 .xls 请先另存为 .xlsx")
    if output_path.exists() and output_path.resolve() != workbook_path.resolve() and not force:
        raise FileExistsError(f"输出文件已存在：{output_path}；确认覆盖请使用 --force")

    workbook = load_workbook(workbook_path)
    if len(workbook.sheetnames) != 1:
        raise ValueError(f"统计表应只有一个工作表，当前为：{workbook.sheetnames}")
    ws = workbook[workbook.sheetnames[0]]
    _, reference_row, detected_start_row = locate_rows(ws)
    start_row = start_row_override or detected_start_row
    if start_row <= reference_row:
        raise ValueError(f"起始行必须大于参考数据行 {reference_row}")

    target_rows = range(start_row, start_row + len(records))
    conflicts: list[str] = []
    for row in target_rows:
        occupied = [
            column
            for column in TARGET_COLUMNS
            if ws[f"{column}{row}"].value not in (None, "")
        ]
        if occupied:
            conflicts.append(f"第 {row} 行（{','.join(occupied)}）")
    if conflicts and not force:
        raise ValueError(
            "目标区域已有数据：" + "；".join(conflicts) + "。确认覆盖请使用 --force"
        )

    for offset, record in enumerate(records):
        row = start_row + offset
        copy_reference_format(ws, reference_row, row)
        for column in TARGET_COLUMNS:
            ws[f"{column}{row}"] = None
        if isinstance(record, FailedDocument):
            ws[f"AN{row}"] = record.note
            ws[f"AN{row}"].number_format = "@"
            continue

        ws[f"U{row}"] = record.hedge_method
        ws[f"V{row}"] = float(record.entry_price)
        ws[f"W{row}"] = float(record.nominal_quantity)
        ws[f"X{row}"] = f"=ROUND(V{row}*W{row},2)"
        ws[f"Y{row}"] = float(record.premium_total)
        ws[f"Z{row}"] = record.otc_option_type
        ws[f"AA{row}"] = float(record.hedge_ratio)
        ws[f"AB{row}"] = record.contract
        ws[f"AC{row}"] = record.signing_date
        ws[f"AD{row}"] = record.effective_date
        ws[f"AE{row}"] = record.expiry_date
        ws[f"BB{row}"] = (
            float(record.exercise_price) if record.exercise_price is not None else None
        )
        ws[f"BN{row}"] = record.pricing_description
        ws[f"AN{row}"] = record.transaction_id

        for column in ("V", "W", "X", "Y", "BB"):
            ws[f"{column}{row}"].number_format = "0.00"
        for column in ("U", "Z", "AB", "AN", "BN"):
            ws[f"{column}{row}"].number_format = "@"
        ws[f"AA{row}"].number_format = "0.00%"
        for column in ("AC", "AD", "AE"):
            ws[f"{column}{row}"].number_format = "yyyy-mm-dd"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}_", suffix=".xlsx", dir=output_path.parent
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return ws.title, start_row, start_row + len(records) - 1


def write_settlement_workbook(
    workbook_path: Path,
    output_path: Path,
    records: Sequence[SettlementRecord],
    start_row_override: int | None = None,
    force: bool = False,
) -> tuple[str, int, int]:
    if workbook_path.suffix.lower() != ".xlsx":
        raise ValueError("当前统计表必须是 .xlsx；旧版 .xls 请先另存为 .xlsx")
    if output_path.exists() and output_path.resolve() != workbook_path.resolve() and not force:
        raise FileExistsError(f"输出文件已存在：{output_path}；确认覆盖请使用 --force")

    workbook = load_workbook(workbook_path)
    if len(workbook.sheetnames) != 1:
        raise ValueError(f"统计表应只有一个工作表，当前为：{workbook.sheetnames}")
    ws = workbook[workbook.sheetnames[0]]
    _, reference_row, detected_start_row = locate_rows_for_columns(
        ws, SETTLEMENT_COLUMNS, "AO"
    )
    start_row = start_row_override or detected_start_row
    if start_row <= reference_row:
        raise ValueError(f"起始行必须大于参考数据行 {reference_row}")

    conflicts: list[str] = []
    for row in range(start_row, start_row + len(records)):
        occupied = [
            column
            for column in SETTLEMENT_COLUMNS
            if ws[f"{column}{row}"].value not in (None, "")
        ]
        if occupied:
            conflicts.append(f"第 {row} 行（{','.join(occupied)}）")
    if conflicts and not force:
        raise ValueError(
            "目标区域已有数据：" + "；".join(conflicts) + "。确认覆盖请使用 --force"
        )

    for offset, record in enumerate(records):
        row = start_row + offset
        copy_reference_format(ws, reference_row, row, SETTLEMENT_COLUMNS)
        for column in SETTLEMENT_COLUMNS:
            ws[f"{column}{row}"] = None
        if isinstance(record, FailedDocument):
            ws[f"AO{row}"] = record.note
            ws[f"AO{row}"].number_format = "@"
            continue

        ws[f"AO{row}"] = record.settlement_id
        ws[f"BB{row}"] = (
            float(record.exercise_price) if record.exercise_price is not None else None
        )
        ws[f"BN{row}"] = record.pricing_description
        ws[f"BF{row}"] = "是" if record.has_payout else "否"
        ws[f"BH{row}"] = float(record.payout_amount)
        ws[f"BJ{row}"] = float(record.settlement_price)

        for column in ("AO", "BF", "BN"):
            ws[f"{column}{row}"].number_format = "@"
        for column in ("BB", "BH", "BJ"):
            ws[f"{column}{row}"].number_format = "0.00"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}_", suffix=".xlsx", dir=output_path.parent
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return ws.title, start_row, start_row + len(records) - 1


def print_preview(records: Sequence[ConfirmationRecord], as_json: bool = False) -> None:
    columns = [
        "文件", "交易编号", "U", "V", "W", "X", "Y", "Z", "AA", "AB",
        "AC", "AD", "AE", "BB", "BN", "AN",
    ]
    rows: list[dict[str, str]] = []
    for record in records:
        if isinstance(record, FailedDocument):
            row = {column: "" for column in columns}
            row["文件"] = record.source_file
            row["交易编号"] = "识别异常"
            row["AN"] = record.note
            rows.append(row)
        else:
            rows.append(record.preview_dict())
    if as_json:
        print(json.dumps(rows, ensure_ascii=False, indent=2))
        return
    print("\t".join(columns))
    for row in rows:
        print("\t".join(row[column] for column in columns))


def print_settlement_preview(
    records: Sequence[SettlementRecord], as_json: bool = False
) -> None:
    columns = ["文件", "AO", "BB", "BN", "BF", "BH", "BJ"]
    rows: list[dict[str, str]] = []
    for record in records:
        if isinstance(record, FailedDocument):
            row = {column: "" for column in columns}
            row["文件"] = record.source_file
            row["AO"] = record.note
            rows.append(row)
        else:
            rows.append(record.preview_dict())
    if as_json:
        print(json.dumps(rows, ensure_ascii=False, indent=2))
        return
    print("\t".join(columns))
    for row in rows:
        print("\t".join(row[column] for column in columns))


def self_test(folder: Path) -> None:
    template = folder / "模板确认书.pdf"
    if not template.exists():
        raise FileNotFoundError(f"未找到测试模板：{template}")
    with PDFTextExtractor() as extractor:
        template_text = extractor.extract(template)
        actual = parse_confirmation(template_text, template.name)
    expected = {
        "transaction_id": "【HFSY】0147-JY-2026082701",
        "entry_price": Decimal("12505.00"),
        "nominal_quantity": Decimal("192.00"),
        "nominal_principal": Decimal("2400960.00"),
        "premium_total": Decimal("101560.32"),
        "otc_option_type": "增强亚式看跌",
        "contract": "LH2701",
        "signing_date": date(2026, 8, 27),
        "effective_date": date(2026, 8, 28),
        "expiry_date": date(2026, 10, 28),
    }
    for field, expected_value in expected.items():
        actual_value = getattr(actual, field)
        if actual_value != expected_value:
            raise AssertionError(f"模板校验失败：{field}={actual_value!r}，应为 {expected_value!r}")
    joined_header_text = template_text.replace("\n签订时间", "签订时间", 1)
    joined_header = parse_confirmation(joined_header_text, "编号与签订时间粘连测试.pdf")
    if joined_header.transaction_id != expected["transaction_id"]:
        raise AssertionError(
            f"交易编号汉字截断失败：{joined_header.transaction_id!r}"
        )
    joined_option_text = template_text.replace("【看跌】\n3.2", "【看跌】 3.2", 1).replace(
        "【增强亚式】\n4、", "【亚式】 4、", 1
    )
    joined_option = parse_confirmation(joined_option_text, "期权字段与后续章节粘连测试.pdf")
    if joined_option.otc_option_type != "亚式看跌":
        raise AssertionError(f"期权字段章节截断失败：{joined_option.otc_option_type!r}")
    if template in discover_pdfs(folder):
        raise AssertionError("生产文件发现逻辑没有排除模板确认书")
    egg = parse_confirmation(template_text.replace("LH2701", "JD2701"), "鸡蛋分支测试.pdf")
    if egg.entry_price != Decimal("25010.00") or egg.nominal_principal != Decimal("4801920.00"):
        raise AssertionError("鸡蛋入场价乘以 2 的规则校验失败")
    print("模板字段、交易编号汉字截断、期权字段章节截断、鸡蛋价格乘 2 及生产扫描排除模板的校验均通过。")


def find_workbook(folder: Path, explicit: str | None) -> Path:
    if explicit:
        path = Path(explicit).expanduser()
        if not path.is_absolute():
            path = folder / path
        return path.resolve()
    candidates = sorted(
        path for path in folder.glob("*.xlsx") if "_已填充" not in path.stem and not path.name.startswith("~$")
    )
    if len(candidates) != 1:
        raise ValueError(f"应找到唯一的原始 .xlsx 统计表，实际为：{[p.name for p in candidates]}")
    return candidates[0].resolve()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="A：处理确认书；B：处理结算单")
    parser.add_argument(
        "function",
        nargs="?",
        type=str.upper,
        choices=("A", "B"),
        help="A=确认书数据提取和写表，B=结算单数据提取和写表",
    )
    parser.add_argument("--folder", default=None, help="覆盖默认数据目录")
    parser.add_argument("--workbook", default=None, help="统计表路径；默认自动寻找唯一原始 .xlsx")
    parser.add_argument("--output", default=None, help="输出路径；默认生成 *_已填充.xlsx")
    parser.add_argument("--preview", action="store_true", help="只提取并打印，不写统计表")
    parser.add_argument("--json", action="store_true", help="预览时输出 JSON")
    parser.add_argument("--force", action="store_true", help="允许覆盖输出文件或目标区域已有数据")
    parser.add_argument("--start-row", type=int, default=None, help="覆盖自动定位的写入起始行")
    parser.add_argument("--self-test", action="store_true", help="仅使用模板确认书运行字段定位测试")
    return parser


def choose_function(function: str | None) -> str:
    if function:
        return function
    if not sys.stdin.isatty():
        raise ValueError("请选择功能：在命令后输入 A（确认书）或 B（结算单）")
    print("请选择功能：")
    print("  A - 确认书数据提取和统计表输入")
    print("  B - 结算单数据提取和统计表输入")
    choice = input("请输入 A 或 B：").strip().upper()
    if choice not in ("A", "B"):
        raise ValueError(f"无效选择：{choice!r}；请输入 A 或 B")
    return choice


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    script_folder = Path(__file__).resolve().parent
    if args.self_test:
        self_test(Path(args.folder).expanduser().resolve() if args.folder else script_folder)
        return 0

    function = choose_function(args.function)
    default_subfolder = "确认书" if function == "A" else "结算单"
    folder = (
        Path(args.folder).expanduser().resolve()
        if args.folder
        else script_folder / default_subfolder
    )
    if not folder.is_dir():
        raise FileNotFoundError(f"未找到{default_subfolder}文件夹：{folder}")

    pdfs = discover_pdfs(folder)
    if not pdfs:
        document_name = "确认书" if function == "A" else "结算单"
        raise FileNotFoundError(f"{folder} 中没有找到{document_name} PDF")

    if function == "A":
        confirmation_records = extract_all(pdfs)
        if args.preview:
            print_preview(confirmation_records, as_json=args.json)
            return 0
    else:
        settlement_records = extract_all_settlements(pdfs)
        if args.preview:
            print_settlement_preview(settlement_records, as_json=args.json)
            return 0

    workbook_path = find_workbook(folder, args.workbook)
    if args.output:
        output_path = Path(args.output).expanduser()
        if not output_path.is_absolute():
            output_path = folder / output_path
        output_path = output_path.resolve()
    else:
        output_path = workbook_path.with_name(f"{workbook_path.stem}_已填充.xlsx")

    if function == "A":
        sheet_name, first_row, last_row = write_workbook(
            workbook_path,
            output_path,
            confirmation_records,
            start_row_override=args.start_row,
            force=args.force,
        )
        print(f"完成：{len(confirmation_records)} 份确认书已写入 {output_path}")
        print(
            f"工作表：{sheet_name}；范围：U{first_row}:AE{last_row}、"
            f"AN{first_row}:AN{last_row}、BB{first_row}:BB{last_row}、"
            f"BN{first_row}:BN{last_row}。"
        )
    else:
        sheet_name, first_row, last_row = write_settlement_workbook(
            workbook_path,
            output_path,
            settlement_records,
            start_row_override=args.start_row,
            force=args.force,
        )
        print(f"完成：{len(settlement_records)} 份结算单已写入 {output_path}")
        print(
            f"工作表：{sheet_name}；范围：AO{first_row}:AO{last_row}、"
            f"BF{first_row}:BF{last_row}、BH{first_row}:BH{last_row}、"
            f"BJ{first_row}:BJ{last_row}、BB{first_row}:BB{last_row}、"
            f"BN{first_row}:BN{last_row}。"
        )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"错误：{exc}", file=sys.stderr)
        raise SystemExit(1)
