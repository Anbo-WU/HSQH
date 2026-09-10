"""规则与真实保单回归测试；输出全部放在临时目录。"""
from contextlib import redirect_stdout
from copy import deepcopy
from decimal import Decimal
from datetime import datetime
import importlib
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import pymupdf

p = importlib.import_module('保单填充')
BASE = Path(__file__).resolve().parent
PDFS = sorted((BASE / '保单0908').glob('*.pdf'))


def tables():
    return [
        [['项目', '保险金额（元）', '保险费\n率（%）', '保险费（元）'],
         ['育肥猪', '10000', '10.0000', '1000']],
        [['条款名称', '交付单位', '其他', '农户', '县（区）财政', '市财政', '省财政', '中央财政'],
         ['条款', '补贴或交付比例（%）', '25', '25', '25', '25', '0', '0'],
         [None, '补贴或交付金额（元）', '250', '250', '250', '250', '0', '0'],
         [None, '补贴或交付金额（元）', '', '', '', '', '', '']],
    ]


class RulesTest(unittest.TestCase):
    def test_location_and_insurance_period(self):
        for address, expected in (
            ('山东省烟台市牟平区姜格庄街道北松', '山东省烟台市牟平区'),
            ('山东省烟台市海阳县某村', '山东省烟台市海阳县'),
            ('广西壮族自治区南宁市武鸣区某镇', '广西壮族自治区南宁市武鸣区'),
        ):
            self.assertEqual(p.district_location('标的地点及方位：\n地点：' + address), expected)
        self.assertEqual(p.coverage_dates('保险期间：自2026年08月28日23:00零时起，至 2027年01月28日二十四时止。'),
                         (datetime(2026, 8, 28), datetime(2027, 1, 28)))
        for text in ('保险期间：自2026年02月30日起，至2027年01月28日止',
                     '保险期间：自2027年01月28日起，至2026年08月28日止',
                     '保险期间：自2026年08月28日起'):
            with self.assertRaises(p.RecognitionError):
                p.coverage_dates(text)
        with self.assertRaises(p.RecognitionError):
            p.district_location('标的地点及方位：地点：山东省烟台市 条款名称：某保险')

    def test_identifiers(self):
        for kind in ('JY', 'FWJY', 'AB', 'ABCD'):
            tid = f'【HFSY】0147-{kind}-2026082705'
            self.assertEqual(p.strict_identifier(tid), tid)
        for tid in ('【HFSY】0147-ABC-2026082705', '【HFSY】0147-JY-2026023005',
                    '【HFSY】0147-jy-2026082705', '【HFSY】0147-JY-20260827050',
                    '【HFSY】014-JY-2026082705', '【HFSY】0147-JY-20260827O5'):
            with self.subTest(tid=tid), self.assertRaises(p.RecognitionError):
                p.strict_identifier(tid)

    def test_company_and_farmers(self):
        self.assertEqual(p.subjects('烟台凯丰源生物科技有限公司'), (0, 1, '烟台凯丰源生物科技有限公司'))
        self.assertEqual(p.subjects('甲有限公司、和丰有限公司'), (0, 2, '甲有限公司、和丰有限公司'))
        self.assertEqual(p.subjects('甲有限公司和乙有限公司'), (0, 2, '甲有限公司、乙有限公司'))
        self.assertEqual(p.subjects('张三、李四，王和志'), (3, 0, ''))
        self.assertEqual(p.subjects('张三 张三'), (1, 0, ''))
        for name in ('详见清单', '张三等10户', '幸福农业合作社', '幸福家庭农场', '张三李四王五'):
            with self.subTest(name=name), self.assertRaises(p.RecognitionError):
                p.subjects(name)

    def test_money_never_accepts_ratio_or_missing(self):
        self.assertEqual(p.number('1,234.50'), Decimal('1234.50'))
        for text in ('25%', '', '1,23.50', '12O.50', '-5'):
            with self.subTest(text=text), self.assertRaises(p.RecognitionError):
                p.number(text)

    def test_amount_row_and_header_mapping(self):
        result = p.table_values(tables())
        self.assertEqual(result['BR'], Decimal(250))
        self.assertEqual(result['BX'], Decimal(250))
        self.assertEqual(result['BT'], Decimal(0))
        self.assertEqual(result['BE'], Decimal('0.1'))

    def test_incomplete_funding_rejected(self):
        for bad in ('', '25%', '249'):
            data = tables()
            data[1][2][2] = bad
            with self.subTest(bad=bad), self.assertRaises(p.RecognitionError):
                p.table_values(data)

    def test_multi_item_or_wrong_rate_rejected(self):
        data = tables()
        data[0].append(['育肥猪', '10000', '10', '1000'])
        with self.assertRaises(p.RecognitionError):
            p.table_values(data)
        data = tables()
        data[0][1][2] = '1'
        with self.assertRaises(p.RecognitionError):
            p.table_values(data)

    def test_batch_does_not_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first = p.output_path(root, '20260908', None)
            first.touch()
            self.assertEqual(p.output_path(root, '20260908', None).name, '已填充数据表2026090802.xlsx')
            with self.assertRaises(FileExistsError):
                p.output_path(root, '20260908', 1)

    def test_workbook_row_updates_duplicates_and_preservation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, out = root / 'source.xlsx', root / 'result.xlsx'
            folder = root / 'pdf'
            folder.mkdir()
            (folder / 'one.pdf').touch()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws['AN4'] = '交易确认书编号'
            ws['AN5'] = '【HFSY】0147-JY-2026082705'
            ws['AN6'] = '【HFSY】0147-JY-2026082706'
            ws['A5'], ws['BA6'], ws['AU5'] = '=1+2', 777, '旧公司名'
            ws['M5'] = 'M列保留'
            wb.save(source)
            original = source.read_bytes()
            values = {**p.table_values(tables()), 'AQ': 2, 'AT': 0, 'AU': None,
                      'BS': Decimal(0), 'BY': Decimal(0), 'BZ': '无',
                      'K': '山东省烟台市牟平区', 'N': datetime(2026, 8, 28), 'O': datetime(2027, 1, 28)}
            record = {'id': ws['AN5'].value, 'path': str(folder / 'one.pdf'), 'method': 'test',
                      'values': values, 'warnings': []}
            with patch.object(p.Reader, 'extract', return_value=record), redirect_stdout(io.StringIO()):
                self.assertEqual(p.run(source, folder, out), (1, 1))
            result = openpyxl.load_workbook(out).active
            self.assertEqual(result['AQ5'].value, 2)
            self.assertEqual(result['AT5'].value, 0)
            self.assertIsNone(result['AU5'].value)
            self.assertEqual(result['A5'].value, '=1+2')
            self.assertEqual(result['BA6'].value, 777)
            self.assertEqual(result['BE5'].value, 0.1)
            self.assertEqual(result['BE5'].number_format, '0.0000%')
            self.assertEqual(result['K5'].value, '山东省烟台市牟平区')
            self.assertEqual(result['M5'].value, 'M列保留')
            self.assertEqual(result['N5'].value, datetime(2026, 8, 28))
            self.assertEqual(result['O5'].value, datetime(2027, 1, 28))
            self.assertEqual(result['N5'].number_format, 'yyyy-mm-dd')
            self.assertEqual(result['O5'].number_format, 'yyyy-mm-dd')
            self.assertEqual(source.read_bytes(), original)
            with self.assertRaises(FileExistsError):
                p.run(source, folder, source)
            (folder / 'duplicate.pdf').touch()
            with patch.object(p.Reader, 'extract', return_value=deepcopy(record)), redirect_stdout(io.StringIO()):
                self.assertEqual(p.run(source, folder, root / 'duplicate.xlsx'), (0, 2))


@unittest.skipUnless(len(PDFS) == 6, '本地六份保单样例不存在')
class RealPolicyTest(unittest.TestCase):
    def test_six_real_policies(self):
        expected = {
            '197': ('05', '3368124.6', '254630.22', '0.0756'),
            '198': ('06', '1354042.2', '102365.59', '0.0756'),
            '199': ('07', '2553196.8', '170298.23', '0.0667'),
            '200': ('02', '2101201.8', '182174.2', '0.0867'),
            '201': ('04', '3787479', '328374.43', '0.0867'),
            '202': ('03', '1402770', '121620.16', '0.0867'),
        }
        reader = p.Reader()
        for pdf in PDFS:
            with self.subTest(pdf=pdf.name):
                record = reader.extract(pdf)
                self.assertEqual(record['values']['K'], '山东省烟台市牟平区')
                self.assertEqual(record['values']['N'], datetime(2026, 8, 28))
                end = {'197': datetime(2027, 1, 28), '198': datetime(2027, 1, 28),
                       '199': datetime(2026, 12, 28)}.get(pdf.stem[-3:], datetime(2027, 2, 28))
                self.assertEqual(record['values']['O'], end)
                suffix, amount, premium, rate = expected[pdf.stem[-3:]]
                self.assertTrue(record['id'].endswith('20260827' + suffix))
                for col, value in zip(('BA', 'BD', 'BE'), (amount, premium, rate)):
                    self.assertEqual(record['values'][col], Decimal(value))

    def test_actual_scan_ocr_matches_native(self):
        native = p.Reader().extract(PDFS[0])
        with tempfile.TemporaryDirectory() as tmp:
            scanned = Path(tmp) / 'scan.pdf'
            with pymupdf.open(PDFS[0]) as original, pymupdf.open() as doc:
                page = original[0]
                image = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
                new_page = doc.new_page(width=page.rect.width, height=page.rect.height)
                new_page.insert_image(new_page.rect, stream=image.tobytes('png'))
                doc.save(scanned)
            with pymupdf.open(scanned) as doc:
                self.assertEqual(doc[0].get_text(), '')
            ocr = p.Reader().extract(scanned)
            self.assertEqual(ocr['method'], 'OCR')
            self.assertEqual(ocr['id'], native['id'])
            self.assertEqual(ocr['values'], native['values'])


if __name__ == '__main__':
    unittest.main()
