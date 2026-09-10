#!/usr/bin/env python3
"""按 AN 确认书编号匹配保单，校验后另存统计表；运行方法见 README.md。"""
from __future__ import annotations

import argparse
import csv
import io
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pymupdf

BASE = Path(__file__).resolve().parent
CENT = Decimal('0.01')
TARGETS = 'K N O AQ AT AU BA BD BE BR BS BT BU BV BW BX BY BZ'.split()
FUNDING = {
    'BR': ('其他',),
    'BT': ('中央财政', '中央财政补贴'),
    'BU': ('省财政', '省级财政', '省级财政补贴'),
    'BV': ('市财政', '市级财政', '市级财政补贴'),
    'BW': ('县(区)财政', '县财政', '县级财政', '县级财政补贴'),
    'BX': ('农户', '农户自缴', '农户自缴保费'),
}
ID_RE = re.compile(r'【HFSY】\d{4}-(?:[A-Z]{4}|[A-Z]{2})-(\d{8})(\d{2})(?![\dA-Za-z])')


class RecognitionError(ValueError):
    """不能可靠提取时，整行跳过，保留待核对记录。"""


def compact(value: object) -> str:
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(value or '')))


def id_text(value: object) -> str:
    return compact(value).translate(str.maketrans({'[': '【', ']': '】', '－': '-', '—': '-', '–': '-'}))


def identifiers(value: object) -> set[str]:
    result = set()
    for match in ID_RE.finditer(id_text(value)):
        try:
            datetime.strptime(match[1], '%Y%m%d')
        except ValueError:
            continue
        result.add(match[0])
    return result


def strict_identifier(value: object) -> str:
    text = id_text(value)
    if not ID_RE.fullmatch(text) or identifiers(text) != {text}:
        raise RecognitionError(f'确认书编号格式或日期无效：{value}')
    return text


def number(value: object, *, percent: bool = False) -> Decimal:
    text = compact(value).replace('人民币', '').replace('￥', '').replace('¥', '').replace('元', '')
    if '%' in text:
        if not percent or not text.endswith('%'):
            raise RecognitionError(f'金额单元格含百分号，拒绝把比例作为金额：{value}')
        text = text[:-1]
    if not re.fullmatch(r'(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?', text):
        raise RecognitionError(f'不是明确的非负数值：{value!r}')
    return Decimal(text.replace(',', ''))


@dataclass(frozen=True)
class Span:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float
    score: float = 1.0

    @property
    def cx(self) -> float:
        return (self.x0 + self.x1) / 2

    @property
    def cy(self) -> float:
        return (self.y0 + self.y1) / 2


def reading_rows(spans: list[Span]) -> list[list[Span]]:
    groups: list[list[Span]] = []
    for span in sorted(spans, key=lambda s: (s.cy, s.x0)):
        if groups and abs(span.cy - sum(s.cy for s in groups[-1]) / len(groups[-1])) < max(
            2, min(span.y1 - span.y0, groups[-1][0].y1 - groups[-1][0].y0) * 0.5
        ):
            groups[-1].append(span)
        else:
            groups.append([span])
    return [sorted(row, key=lambda s: s.x0) for row in groups]


def region_text(spans: list[Span]) -> str:
    return '\n'.join(' '.join(s.text for s in row) for row in reading_rows(spans))


def anchor_y(spans: list[Span], label: str) -> float:
    rows = [row for row in reading_rows(spans) if label in compact(''.join(s.text for s in row))]
    if len(rows) != 1:
        raise RecognitionError(f'无法唯一定位“{label}”区域')
    return min(s.y0 for s in rows[0])


def subjects(value: str) -> tuple[int, int, str]:
    """仅解析姓名栏；不把地址、投保人、捐助公司或附件说明当成被保险人。"""
    value = unicodedata.normalize('NFKC', value).strip(' :：;；')
    if not value or re.search(r'详见|见附|清单|等\d*户|等人', compact(value)):
        raise RecognitionError(f'姓名栏不完整，需要核对农户名单：{value}')
    if '有限公司' in compact(value):
        text = re.sub(r'(?<=有限公司)[和及与]', '、', compact(value))
        companies = re.findall(r'[\u4e00-\u9fffA-Za-z0-9()]+?有限公司', text)
        remainder = text
        for company in companies:
            remainder = remainder.replace(company, '', 1)
        if not companies or remainder.strip('、,;和及与'):
            raise RecognitionError(f'公司姓名栏存在未识别内容：{value}')
        companies = list(dict.fromkeys(companies))
        return 0, len(companies), '、'.join(companies)
    # 没有“有限公司”也不能把合作社、家庭农场等组织误计成自然人。
    if re.search(r'公司|合作社|农场|养殖场|中心|协会|企业|经营部|有限合伙', value):
        raise RecognitionError(f'主体既非有限公司，也不是可确认的自然人名单：{value}')
    names = [n for n in re.split(r'[\s、,;；/]+', value) if n]
    if not names or any(not re.fullmatch(r'[\u4e00-\u9fff]{2,4}|[\u4e00-\u9fff]{2,12}·[\u4e00-\u9fff·]{2,15}', n) for n in names):
        raise RecognitionError(f'无法可靠分隔自然人姓名：{value}')
    return len(set(names)), 0, ''


def insured(spans: list[Span]) -> tuple[int, int, str]:
    start, end = anchor_y(spans, '被保险人信息'), anchor_y(spans, '保障内容')
    text = region_text([s for s in spans if start <= s.cy < end])
    match = re.search(r'姓\s*名\s*[/／]\s*单\s*位\s*名\s*称\s*[:：]\s*(.+?)(?=联系电话|证件类型|证件号码|联系地址|邮编|$)', text, re.S)
    if not match:
        raise RecognitionError('未找到被保险人姓名/单位名称')
    return subjects(match[1].strip())


def body_identifier(spans: list[Span]) -> str:
    start, end = anchor_y(spans, '特别约定'), anchor_y(spans, '签单公司信息')
    ids = identifiers(region_text([s for s in spans if start <= s.cy < end]))
    if len(ids) != 1:
        raise RecognitionError(f'第一页特别约定中未找到唯一有效确认书编号：{sorted(ids)}')
    return next(iter(ids))


def clean_header(value: object) -> str:
    return re.sub(r'\((?:元|%)\)', '', compact(value))


def district_location(text: str) -> str:
    text = compact(text)
    match = re.search(r'标的地点及方位[:：]?地点[:：]([^:：;；]+)', text)
    if not match:
        raise RecognitionError('未找到标的地点及方位下的“地点：”')
    address = re.split(r'条款名称|保险期间|方位[:：]', match[1])[0]
    for ending in re.finditer(r'[区县]', address):
        location = address[:ending.end()]
        if location.endswith('自治区'):
            continue
        if re.fullmatch(r'[\u4e00-\u9fff·]+', location):
            return location
    raise RecognitionError('标的地点未识别到明确的区/县')


def coverage_dates(text: str) -> tuple[datetime, datetime]:
    text = compact(text)
    date_part = r'(\d{4}年\d{1,2}月\d{1,2}日)'
    matches = list(re.finditer(r'保险期间[:：]自' + date_part + r'[^至]*至' + date_part, text))
    if len(matches) != 1:
        raise RecognitionError('未找到唯一完整的保险期间起止日期')
    try:
        start, end = (datetime.strptime(value, '%Y年%m月%d日') for value in matches[0].groups())
    except ValueError as exc:
        raise RecognitionError('保险期间包含无效日期') from exc
    if end < start:
        raise RecognitionError('保险到期日早于起始日')
    return start, end


def policy_details(spans: list[Span]) -> dict:
    location_y = anchor_y(spans, '标的地点及方位')
    period_y = anchor_y(spans, '保险期间')
    end_y = anchor_y(spans, '特别约定')
    location_spans = [s for s in spans if location_y <= s.cy < period_y]
    period_spans = [s for s in spans if period_y <= s.cy < end_y]
    # 新字段的 OCR 低置信度也沿用整行转人工核对规则。
    location_rows = reading_rows(location_spans)
    relevant_location = []
    for row in location_rows:
        if '条款名称' in compact(''.join(s.text for s in row)):
            break
        relevant_location.extend(row)
    if any(s.score < 0.85 for s in relevant_location + period_spans):
        raise RecognitionError('标的地点或保险期间 OCR 置信度不足')
    start, end = coverage_dates(region_text(period_spans))
    return {'K': district_location(region_text(relevant_location)), 'N': start, 'O': end}


def table_values(tables: list[list[list[str | None]]]) -> dict[str, Decimal]:
    coverage = []
    funding = []
    for table in tables:
        for ri, row in enumerate(table):
            headers = [clean_header(c) for c in row]
            if all(label in headers for label in ('保险金额', '保险费率', '保险费')):
                coverage.append((table, ri, headers))
            if all(any(clean_header(a) in headers for a in aliases) for aliases in FUNDING.values()):
                funding.append((table, ri, headers))
    if len(coverage) != 1 or len(funding) != 1:
        raise RecognitionError('未唯一定位保障内容表及保险费构成表（或表头识别不完整）')
    table, ri, headers = coverage[0]
    cols = [headers.index(label) for label in ('保险金额', '保险费率', '保险费')]
    entries = []
    for row in table[ri + 1:]:
        cells = [row[c] if c < len(row) else '' for c in cols]
        if not any(compact(c) for c in cells):
            continue
        entries.append(tuple(number(c, percent=i == 1) for i, c in enumerate(cells)))
    if len(entries) != 1:
        raise RecognitionError(f'保障内容有 {len(entries)} 条金额记录，需要核对多标的/合计，未自动合并')
    amount, rate, premium = entries[0]
    if amount <= 0 or not 0 <= rate <= 100:
        raise RecognitionError('保险金额或百分比费率超出有效范围')
    # 保单四舍五入到分；不使用公式反推值覆盖原始费率。
    if abs(amount * rate / 100 - premium) > CENT:
        raise RecognitionError(f'保险金额×费率与保费不一致：{amount} × {rate}% ≠ {premium}')
    values = {'BA': amount, 'BD': premium, 'BE': rate / 100}
    table, ri, headers = funding[0]
    indexes = {col: next(i for i, h in enumerate(headers) if h in tuple(map(clean_header, aliases)))
               for col, aliases in FUNDING.items()}
    records = []
    for row in table[ri + 1:]:
        if not any('金额' in compact(c) and '比例' not in compact(c) for c in row):
            continue
        cells = {col: row[i] if i < len(row) else '' for col, i in indexes.items()}
        if not any(compact(c) for c in cells.values()):
            continue
        # 空白不是零，避免把漏识别数字当作无补贴。
        records.append({col: number(c) for col, c in cells.items()})
        # 对模板之外另有资金来源的情况报错，不静默丢弃金额。
        for i, cell in enumerate(row):
            if i not in indexes.values() and i < len(headers) and not headers[i] and compact(cell):
                raise RecognitionError(f'保险费构成出现无表头的额外数据：{cell}')
    if not records:
        raise RecognitionError('保险费构成中没有完整的金额行')
    values.update({col: sum((record[col] for record in records), Decimal(0)) for col in FUNDING})
    if abs(sum(values[col] for col in FUNDING) - premium) > CENT:
        raise RecognitionError('保险费构成金额合计与保费总额不一致')
    return values


class Reader:
    def __init__(self, force_ocr: bool = False):
        self.engine = None
        self.force_ocr = force_ocr

    def ocr_page(self, page) -> tuple[list[Span], list[list[list[str]]]]:
        import cv2
        import numpy as np
        from rapidocr import RapidOCR

        if self.engine is None:
            self.engine = RapidOCR()
        scale = 3
        pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), colorspace=pymupdf.csRGB, alpha=False)
        rgb = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, 3)
        # RapidOCR ndarray 输入采用 BGR。
        bgr = cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)
        result = self.engine(bgr)
        if result.txts is None or result.boxes is None or result.scores is None:
            raise RecognitionError('OCR 未返回文字、坐标和置信度')
        pixel_spans = []
        for text, box, score in zip(result.txts, result.boxes, result.scores, strict=True):
            pixel_spans.append(Span(text, float(box[:, 0].min()), float(box[:, 1].min()),
                                    float(box[:, 0].max()), float(box[:, 1].max()), float(score)))
        spans = [Span(s.text, s.x0 / scale, s.y0 / scale, s.x1 / scale, s.y1 / scale, s.score)
                 for s in pixel_spans]
        # 关键正文低置信度直接转人工核对，避免数字/姓名被静默误填。
        relevant = ('被保险人信息', '保障内容', '特别约定', '签单公司信息')
        bounds = [anchor_y(spans, label) for label in relevant]
        for s in spans:
            if (bounds[0] <= s.cy < bounds[1] or bounds[2] <= s.cy < bounds[3]) and s.score < 0.85:
                raise RecognitionError(f'OCR 关键正文置信度不足：{s.text} ({s.score:.2f})')

        gray = cv2.cvtColor(bgr, cv2.COLOR_BGR2GRAY)
        binary = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)[1]
        horizontal = cv2.morphologyEx(binary, cv2.MORPH_OPEN, np.ones((1, max(25, pix.width // 35)), np.uint8))
        vertical = cv2.morphologyEx(binary, cv2.MORPH_OPEN, np.ones((max(25, pix.height // 70), 1), np.uint8))
        grid = cv2.bitwise_or(horizontal, vertical)
        contours, _ = cv2.findContours(grid, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

        def positions(mask, axis, minimum):
            indices = np.where(np.count_nonzero(mask, axis=axis) > minimum)[0]
            groups = np.split(indices, np.where(np.diff(indices) > 3)[0] + 1)
            return [int(np.mean(g)) for g in groups if len(g)]

        tables = []
        for contour in contours:
            x, y, w, h = cv2.boundingRect(contour)
            if w < pix.width * 0.5 or h < pix.height * 0.06:
                continue
            xs = [x + v for v in positions(vertical[y:y+h, x:x+w], 0, h * 0.45)]
            ys = [y + v for v in positions(horizontal[y:y+h, x:x+w], 1, w * 0.65)]
            if len(xs) < 5 or len(ys) < 3:
                continue
            table = []
            for top, bottom in zip(ys, ys[1:]):
                row = []
                for left, right in zip(xs, xs[1:]):
                    cell = [s for s in pixel_spans if left <= s.cx < right and top <= s.cy < bottom]
                    if any(s.score < 0.85 for s in cell):
                        raise RecognitionError('OCR 表格单元格置信度不足，请检查扫描清晰度')
                    row.append(region_text(cell))
                table.append(row)
            tables.append(table)
        return spans, tables

    def extract(self, path: Path) -> dict:
        with pymupdf.open(path) as doc:
            if not len(doc):
                raise RecognitionError('PDF 没有页面')
            page = doc[0]
            if not self.force_ocr:
                try:
                    spans = [Span(w[4], *w[:4]) for w in page.get_text('words')]
                    tables = [t.extract() for t in page.find_tables().tables]
                    return self.parse(path, spans, tables, '文字层')
                except (RecognitionError, ValueError, RuntimeError) as exc:
                    print(f'  文字层未通过校验，改用 OCR：{exc}', flush=True)
            print(f'  OCR：{path.name} 第 1 页', flush=True)
            spans, tables = self.ocr_page(page)
            return self.parse(path, spans, tables, 'OCR')

    @staticmethod
    def parse(path: Path, spans, tables, method: str) -> dict:
        transaction_id = body_identifier(spans)
        file_ids = identifiers(path.stem)
        if file_ids and file_ids != {transaction_id}:
            raise RecognitionError(f'文件名与特别约定编号不一致：{file_ids} / {transaction_id}')
        farmers, companies, names = insured(spans)
        values = table_values(tables)
        values.update(policy_details(spans))
        values.update({'AQ': farmers, 'AT': companies, 'AU': names or None,
                       'BS': Decimal(0), 'BY': Decimal(0), 'BZ': '无'})
        warnings = [f'异常：同一保单有 {companies} 家公司：{names}'] if companies > 1 else []
        return {'id': transaction_id, 'path': str(path), 'method': method, 'values': values, 'warnings': warnings}


def output_path(folder: Path, day: str, batch: int | None) -> Path:
    if batch is not None and batch < 1:
        raise ValueError('批次必须是正整数')
    current = batch or 1
    while True:
        path = folder / f'已填充数据表{day}{current:02d}.xlsx'
        if not path.exists() and not path.with_suffix('.csv').exists():
            return path
        if batch is not None:
            raise FileExistsError(f'该批次结果已存在，请换一个批次：{path}')
        current += 1


def run(source: Path, folder: Path, output: Path | None, *, sheet: str | None = None,
        force_ocr: bool = False) -> tuple[int, int]:
    if not folder.is_dir():
        raise FileNotFoundError(f'未找到保单目录：{folder}')
    pdfs = sorted(p for p in folder.rglob('*') if p.suffix.lower() == '.pdf')
    if not pdfs:
        raise FileNotFoundError(f'目录中没有 PDF：{folder}')
    if source.suffix.lower() != '.xlsx':
        raise ValueError('统计表请使用 .xlsx 格式')
    if output and (source.resolve() == output.resolve() or output.exists() or output.with_suffix('.csv').exists()):
        raise FileExistsError('输出必须是尚不存在的新文件，不能覆盖源表或已有结果')
    book = openpyxl.load_workbook(source)
    if sheet and sheet not in book.sheetnames:
        raise ValueError(f'工作表不存在：{sheet}')
    selected = [book[sheet]] if sheet else book.worksheets
    jobs = []
    for ws in selected:
        headers = [row for row in range(1, min(ws.max_row, 30) + 1)
                   if '确认书编号' in compact(ws[f'AN{row}'].value)]
        if not headers:
            continue
        for row in range(headers[0] + 1, ws.max_row + 1):
            if ws[f'AN{row}'].value is not None and compact(ws[f'AN{row}'].value):
                jobs.append((ws, row, ws[f'AN{row}'].value))
    if not jobs:
        raise ValueError('未找到 AN 列的确认书编号数据（应有“交易确认书编号”表头）')
    reader = Reader(force_ocr)
    index: dict[str, list[dict]] = {}
    report = []
    for path in pdfs:
        print(f'读取：{path.name}', flush=True)
        try:
            record = reader.extract(path)
            index.setdefault(record['id'], []).append(record)
        except Exception as exc:
            message = f'{type(exc).__name__}: {exc}'
            print(f'  [识别失败] {message}', flush=True)
            report.append({'状态': 'PDF识别失败', 'PDF': str(path), '说明': message})
    success = 0
    for ws, row, raw_id in jobs:
        item = {'工作表': ws.title, '行号': row, '确认书编号': raw_id}
        try:
            tid = strict_identifier(raw_id)
            matches = index.get(tid, [])
            if not matches:
                raise RecognitionError('没有找到正文编号匹配且全部校验通过的保单')
            if len(matches) > 1:
                raise RecognitionError('多份保单对应同一编号：' + '、'.join(Path(m['path']).name for m in matches))
            record = matches[0]
            for col in TARGETS:
                if isinstance(ws[f'{col}{row}'], openpyxl.cell.cell.MergedCell):
                    raise RecognitionError(f'{col}{row} 是合并单元格，不能填充')
            for col, value in record['values'].items():
                cell = ws[f'{col}{row}']
                cell.value = float(value) if isinstance(value, Decimal) else value
                cell.number_format = ('yyyy-mm-dd' if col in ('N', 'O') else
                                      '0.0000%' if col == 'BE' else
                                      '0' if col in ('AQ', 'AT') else
                                      'General' if col in ('K', 'AU', 'BZ') else '#,##0.00')
            success += 1
            item.update({'状态': '已填充' if output else '预览通过', 'PDF': record['path'], '识别方式': record['method'],
                         '说明': '；'.join(record['warnings']), **record['values']})
            print(f'  [通过] {ws.title} 第 {row} 行 {tid}：{record["values"]["AU"] or str(record["values"]["AQ"]) + " 位自然人"}', flush=True)
            for warning in record['warnings']:
                print(f'  [异常] {warning}', flush=True)
        except RecognitionError as exc:
            item.update({'状态': '未填充', '说明': str(exc)})
            print(f'  [未填充] {ws.title} 第 {row} 行 {raw_id}：{exc}', flush=True)
        report.append(item)
    used = {id_text(raw_id) for _, _, raw_id in jobs}
    for tid, records in index.items():
        if tid not in used:
            for record in records:
                report.append({'状态': '保单无对应统计行', '确认书编号': tid, 'PDF': record['path']})
    if output:
        # 先在内存生成完整工作簿，再以独占创建模式落盘，避免覆盖已有结果。
        stream = io.BytesIO()
        book.save(stream)
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open('xb') as handle:
            handle.write(stream.getvalue())
        fields = ['状态', '工作表', '行号', '确认书编号', 'PDF', '识别方式', '说明', *TARGETS]
        with output.with_suffix('.csv').open('x', encoding='utf-8-sig', newline='') as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows({key: value.strftime('%Y-%m-%d') if isinstance(value, datetime) else value
                              for key, value in item.items()} for item in report)
        print(f'新表：{output}\n核对记录：{output.with_suffix(".csv")}', flush=True)
    book.close()
    skipped = len(jobs) - success
    print(f'{"完成" if output else "预览"}：成功 {success} 行，未填充 {skipped} 行。源统计表未修改。', flush=True)
    return success, skipped


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=BASE / '保期数据统计表.xlsx', help='源统计表 .xlsx')
    parser.add_argument('--date', default=datetime.now().strftime('%Y%m%d'), help='业务日期 YYYYMMDD，默认今天')
    parser.add_argument('--pdf-dir', type=Path, help='保单目录；默认脚本旁 保单MMDD 或 保单YYYYMMDD')
    parser.add_argument('--output-dir', type=Path, default=BASE, help='新表和核对记录输出目录')
    parser.add_argument('--batch', type=int, help='批次正整数；不填则自动取当天未使用的下一批次')
    parser.add_argument('--sheet', help='只处理指定工作表；默认处理有 AN 表头的所有工作表')
    parser.add_argument('--preview', action='store_true', help='只识别和校验，不生成文件')
    parser.add_argument('--ocr', action='store_true', help='强制使用本地 OCR，可用于检查扫描件识别')
    args = parser.parse_args(argv)
    try:
        if not re.fullmatch(r'\d{8}', args.date):
            raise ValueError('--date 必须是 YYYYMMDD')
        datetime.strptime(args.date, '%Y%m%d')
        folder = args.pdf_dir
        if folder is None:
            choices = [BASE / f'保单{args.date}', BASE / f'保单{args.date[4:]}']
            existing = [p for p in choices if p.is_dir()]
            if len(existing) > 1:
                raise ValueError('同时存在两种日期格式的保单目录，请用 --pdf-dir 指定')
            folder = existing[0] if existing else choices[1]
        output = None if args.preview else output_path(args.output_dir, args.date, args.batch)
        _, skipped = run(args.input, folder, output, sheet=args.sheet, force_ocr=args.ocr)
        return 2 if skipped else 0
    except (OSError, ValueError, KeyError) as exc:
        print(f'错误：{exc}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
