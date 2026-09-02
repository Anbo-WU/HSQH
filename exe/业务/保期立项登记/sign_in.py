#!/usr/bin/env python3
"""遍历所有子文件夹中的立项申请表，并填充业务记录.xlsx。

依赖：python3 -m pip install openpyxl
用法：python3 填充业务记录.py
"""

from __future__ import annotations

import argparse
import re
from calendar import monthrange
from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - 给缺少依赖的环境更明确的提示
    raise SystemExit(
        "缺少 openpyxl，请先运行：python3 -m pip install openpyxl"
    ) from exc


TARGET_NAME = "业务记录.xlsx"
START_SERIAL = 23
OUTPUT_COLUMNS = 15
START_ROW = 3
PERCENT_COLUMNS = (9, 10, 11, 14)  # I、J、K、N，格式以模板第 2 行为准

CHECKED_MARKS = ("☑", "✓", "✔", "√", "■", "●")
UNCHECKED_MARKS = ("□", "☐", "○")


def natural_text_key(text: str) -> tuple[tuple[int, Any], ...]:
    parts = re.split(r"(\d+)", text)
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in parts
    )


def natural_key(path: Path) -> tuple[tuple[int, Any], ...]:
    """让 2.xlsx 排在 10.xlsx 前，并保持同名规则下的稳定顺序。"""
    return natural_text_key(path.name)


def relative_natural_key(path: Path, root: Path) -> tuple[tuple[int, Any], ...]:
    """按相对路径逐级自然排序，保证批次文件夹顺序稳定。"""
    result: list[tuple[int, Any]] = []
    for part in path.relative_to(root).parts:
        result.extend(natural_text_key(part))
    return tuple(result)


def source_files(root: Path) -> list[Path]:
    """只读取 root 的子文件夹，递归收集其中的 Excel 申请表。"""
    files: list[Path] = []
    folders = sorted(
        (path for path in root.rglob("*") if path.is_dir()),
        key=lambda path: relative_natural_key(path, root),
    )
    for folder in folders:
        candidates = (
            p
            for p in folder.iterdir()
            if p.is_file()
            and not p.name.startswith(("~$", ".~"))
            and p.suffix.lower() in {".xlsx", ".xlsm"}
        )
        files.extend(sorted(candidates, key=natural_key))
    return files


def normalized_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value)).replace("（", "(").replace("）", ")")


def row_has_label(ws: Any, row: int, keywords: Iterable[str]) -> bool:
    """在 A、B 两列验证字段名；兼容合并单元格的标签。"""
    label = normalized_text(ws.cell(row, 1).value) + normalized_text(ws.cell(row, 2).value)
    return all(normalized_text(word) in label for word in keywords)


def find_label_row(ws: Any, keyword_groups: Iterable[Iterable[str]]) -> int | None:
    for keywords in keyword_groups:
        for row in range(1, ws.max_row + 1):
            if row_has_label(ws, row, keywords):
                return row
    return None


def value_from_row(ws: Any, row: int, columns: Iterable[int]) -> Any:
    for column in columns:
        value = ws.cell(row, column).value
        if value not in (None, ""):
            return value
    return None


def field_value(
    ws: Any,
    preferred_row: int,
    columns: Iterable[int],
    keyword_groups: Iterable[Iterable[str]],
) -> Any:
    """优先使用用户指定行；标签不匹配时兼容不同版本的模板。"""
    groups = tuple(tuple(group) for group in keyword_groups)
    if any(row_has_label(ws, preferred_row, group) for group in groups):
        return value_from_row(ws, preferred_row, columns)
    actual_row = find_label_row(ws, groups)
    if actual_row is None:
        raise ValueError(f"未找到字段：{' / '.join('+'.join(g) for g in groups)}")
    return value_from_row(ws, actual_row, columns)


def checked_option(values: Iterable[Any]) -> str:
    for value in values:
        text = "" if value is None else str(value).strip()
        if any(mark in text for mark in CHECKED_MARKS):
            for mark in CHECKED_MARKS + UNCHECKED_MARKS:
                text = text.replace(mark, "")
            return text.strip()
    raise ValueError("指定单元格中没有找到打勾的选项")


CHINESE_NUMBERS = {
    "零": 0,
    "一": 1,
    "二": 2,
    "两": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
    "十": 10,
    "十一": 11,
    "十二": 12,
}


def mentions_full_sampling(ws: Any) -> bool:
    """检查采价周期之后的项目说明，避开 G19 中仅用于提示填写的示例文字。"""
    for row in ws.iter_rows(min_row=20):
        for cell in row:
            if cell.value is not None and "全程采价" in normalized_text(cell.value):
                return True
    return False


def normalized_sampling_period(value: Any, ws: Any) -> str:
    """按采价方式输出“最后1个月”“全程采价”或源表中的其他周期。"""
    if value in (None, ""):
        raise ValueError("采价周期为空")

    text = normalized_text(value)
    if "保险到期日前30个交易日" in text:
        return "最后1个月"
    if "最后" in text and ("30个交易日" in text or "一个月" in text or "1个月" in text):
        return "最后1个月"
    if "全程" in text or mentions_full_sampling(ws):
        return "全程采价"

    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(str(value))
        text = format(number.normalize(), "f")
        return f"{text}个月"

    match = re.search(r"(\d+(?:\.\d+)?)个?月", text)
    if match:
        number = Decimal(match.group(1))
        return f"{format(number.normalize(), 'f')}个月"

    match = re.search(r"([零一二两三四五六七八九十]+)个?月", text)
    if match and match.group(1) in CHINESE_NUMBERS:
        return f"{CHINESE_NUMBERS[match.group(1)]}个月"

    raise ValueError(f"无法规范采价周期：{value!r}")


def yes_or_no(value: Any) -> str:
    if value in (None, ""):
        raise ValueError("是否保底字段为空")
    try:
        return "否" if Decimal(str(value).strip()) == 0 else "是"
    except InvalidOperation:
        text = normalized_text(value)
        if text in {"否", "无", "0", "0%"}:
            return "否"
        if text in {"是", "有"}:
            return "是"
        raise ValueError(f"无法判断是否保底：{value!r}")


def add_months(day: date, months: int) -> date:
    month_index = day.year * 12 + day.month - 1 + months
    year, month_zero = divmod(month_index, 12)
    month = month_zero + 1
    return date(year, month, min(day.day, monthrange(year, month)[1]))


def parse_date_range(value: Any) -> tuple[date, date]:
    if isinstance(value, (date, datetime)):
        raise ValueError("产品周期只有一个日期，缺少起止日期")
    text = str(value or "")
    matches = re.findall(r"(20\d{2})\s*[年/.\-]\s*(\d{1,2})\s*[月/.\-]\s*(\d{1,2})\s*日?", text)
    if len(matches) < 2:
        raise ValueError(f"无法从产品周期中识别两个日期：{value!r}")
    start, end = (date(*(int(part) for part in match)) for match in matches[:2])
    if end < start:
        raise ValueError(f"产品周期结束日期早于开始日期：{value!r}")
    return start, end


def rounded_months(value: Any) -> str:
    """按日历整月计算，剩余不足/超过半个月时严格四舍五入。"""
    start, end = parse_date_range(value)
    whole = (end.year - start.year) * 12 + end.month - start.month
    if add_months(start, whole) > end:
        whole -= 1
    anchor = add_months(start, whole)
    next_anchor = add_months(start, whole + 1)
    fraction = Decimal((end - anchor).days) / Decimal((next_anchor - anchor).days)
    months = int((Decimal(whole) + fraction).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return f"{months}个月"


NUMBER_RE = re.compile(r"[-+]?\d[\d,]*(?:\.\d+)?(?:[eE][-+]?\d+)?")


def number_only(value: Any) -> int | float:
    """删除“万元”等文字；整数写成整数，小数保留其有效小数位。"""
    if value is None or value == "":
        raise ValueError("数字字段为空")
    match = NUMBER_RE.search(str(value))
    if not match:
        raise ValueError(f"未找到数字：{value!r}")
    number = Decimal(match.group(0).replace(",", ""))
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def percentage_value(value: Any) -> int | float:
    """百分比以小数数值写入；保留到万分位，对应百分号后的两位小数。"""
    if value is None or value == "":
        raise ValueError("百分比字段为空")
    match = NUMBER_RE.search(str(value))
    if not match:
        raise ValueError(f"未找到百分比数字：{value!r}")
    number = Decimal(match.group(0).replace(",", ""))
    if "%" in str(value):
        number /= Decimal(100)
    number = number.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
    return int(number) if number == number.to_integral_value() else float(number)


def extract_record(path: Path) -> list[Any]:
    # data_only=True 读取公式在 Excel 中最后保存的计算结果。
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        ws = workbook["填写页"] if "填写页" in workbook.sheetnames else workbook.worksheets[0]

        project = field_value(ws, 3, (3, 4, 5), (("项目名称",),))
        variety = field_value(ws, 5, (3, 4, 5), (("项目品种",),))
        period_value = field_value(
            ws, 31, (3, 4, 5), (("保险期限",), ("产品周期",))
        )
        sampling = field_value(ws, 19, (3, 4, 5), (("采价周期",),))
        guarantee = field_value(ws, 18, (3, 4, 5), (("约定赔付金额",), ("保底",)))
        moneyness = field_value(ws, 25, (3, 4, 5), (("虚实值程度",),))
        insurance_rate = field_value(ws, 27, (4, 5), (("预计保险费",), ("保险费率",)))
        option_rate = field_value(ws, 22, (4, 5), (("预计期权费",), ("期权费率",)))
        total_premium = field_value(ws, 27, (3,), (("预计保险费",), ("预计总保费",)))
        support_amount = field_value(ws, 32, (3,), (("申请支持资金金额",),))
        support_ratio = field_value(ws, 32, (4, 5), (("申请支持资金金额",),))
        expected_profit = field_value(ws, 26, (3, 4, 5), (("预期利润",),))

        structure_row = find_label_row(ws, (("期权结构",),)) or 15
        structure = checked_option(ws.cell(structure_row, col).value for col in (3, 4, 5))
        direction_row = find_label_row(ws, (("期权方向",),)) or 14
        direction = checked_option(ws.cell(direction_row, col).value for col in (3, 4, 5))
        option_structure = f"{structure}{direction}"

        folder_match = re.search(r"(\d{2})$", path.parent.name)
        if not folder_match:
            raise ValueError(f"无法从文件夹名称提取批次：{path.parent.name}")
        review_batch = f"{int(folder_match.group(1))}次"

        return [
            None,  # A 列序号稍后统一生成
            project,
            review_batch,
            variety,
            rounded_months(period_value),
            normalized_sampling_period(sampling, ws),
            option_structure,
            yes_or_no(guarantee),
            percentage_value(moneyness),
            percentage_value(insurance_rate),
            percentage_value(option_rate),
            number_only(total_premium),
            number_only(support_amount),
            percentage_value(support_ratio),
            number_only(expected_profit),
        ]
    finally:
        workbook.close()


def copy_row_style(ws: Any, source_row: int, target_row: int) -> None:
    for column in range(1, OUTPUT_COLUMNS + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def write_records(target: Path, records: list[list[Any]], output: Path) -> None:
    workbook = load_workbook(target)
    try:
        ws = workbook.worksheets[0]
        start_row = START_ROW
        style_row = 2

        for offset, record in enumerate(records):
            row = start_row + offset
            record[0] = START_SERIAL + offset
            copy_row_style(ws, style_row, row)
            for column, value in enumerate(record, start=1):
                ws.cell(row, column).value = value
            # 所有百分比列严格沿用第 2 行模板的百分比显示格式。
            for column in PERCENT_COLUMNS:
                ws.cell(row, column).number_format = ws.cell(style_row, column).number_format

        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--root", type=Path, default=Path(__file__).resolve().parent, help="主文件夹路径"
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="输出文件；默认直接更新主文件夹中的业务记录.xlsx",
    )
    parser.add_argument(
        "--preview", action="store_true", help="仅显示将写入的数据，不修改工作簿"
    )
    args = parser.parse_args()

    root = args.root.resolve()
    target = root / TARGET_NAME
    if not target.is_file():
        raise FileNotFoundError(f"找不到目标文件：{target}")

    paths = source_files(root)
    if not paths:
        raise FileNotFoundError("所有子文件夹中都没有找到 .xlsx 或 .xlsm 文件")

    records: list[list[Any]] = []
    for index, path in enumerate(paths, start=START_SERIAL):
        try:
            record = extract_record(path)
        except Exception as exc:
            raise RuntimeError(f"提取失败：{path}") from exc
        record[0] = index
        records.append(record)
        print(f"{index}: {path.relative_to(root)}")
        if args.preview:
            print("   ", record)

    if args.preview:
        print(f"预览完成，共 {len(records)} 条；未修改 {TARGET_NAME}")
        return 0

    output = args.output.resolve() if args.output else target
    if output == target:
        lock_files = (root / f"~${TARGET_NAME}", root / f".~{TARGET_NAME}")
        if any(lock_file.exists() for lock_file in lock_files):
            raise RuntimeError(
                f"检测到 {TARGET_NAME} 可能正被打开，请关闭该工作簿后再运行脚本"
            )
    write_records(target, records, output)
    print(f"完成：已向 {output} 写入 {len(records)} 条记录（序号 {START_SERIAL}-{START_SERIAL + len(records) - 1}）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
