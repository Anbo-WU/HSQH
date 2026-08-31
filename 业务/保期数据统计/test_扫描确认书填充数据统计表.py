#!/usr/bin/env python3
"""模板确认书仅在本测试中作为定位与回归样本。"""

from __future__ import annotations

import tempfile
import unittest
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from openpyxl import load_workbook

import 扫描确认书填充数据统计表 as app


class ConfirmationExtractionTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.extractor = app.PDFTextExtractor()
        cls.template_path = HERE / "模板确认书.pdf"
        cls.template_text = cls.extractor.extract(cls.template_path)
        cls.record = app.parse_confirmation(cls.template_text, cls.template_path.name)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.extractor.close()

    def test_template_fields(self) -> None:
        record = self.record
        self.assertEqual(record.transaction_id, "【HFSY】0147-JY-2026082701")
        self.assertEqual(record.entry_price, Decimal("12505.00"))
        self.assertEqual(record.nominal_quantity, Decimal("192.00"))
        self.assertEqual(record.nominal_principal, Decimal("2400960.00"))
        self.assertEqual(record.premium_total, Decimal("101560.32"))
        self.assertEqual(record.otc_option_type, "增强亚式看跌")
        self.assertEqual(record.contract, "LH2701")
        self.assertEqual(record.signing_date, date(2026, 8, 27))
        self.assertEqual(record.effective_date, date(2026, 8, 28))
        self.assertEqual(record.expiry_date, date(2026, 10, 28))

    def test_production_discovery_excludes_template(self) -> None:
        files = app.discover_pdfs(HERE)
        self.assertNotIn(self.template_path, files)

        with tempfile.TemporaryDirectory(prefix="discovery_test_") as tempdir:
            folder = Path(tempdir)
            expected = [
                folder / "任意名称.pdf",
                folder / "客户_结算通知书_001.pdf",
                folder / "客户_商品交易确认书_001.pdf",
            ]
            for name in (
                *(path.name for path in expected),
                "模板确认书.pdf",
            ):
                (folder / name).write_bytes(b"test")
            self.assertEqual(app.discover_pdfs(folder), sorted(expected, key=lambda path: path.name))

    def test_duplicate_pdf_content_is_processed_once(self) -> None:
        with tempfile.TemporaryDirectory(prefix="duplicate_test_") as tempdir:
            folder = Path(tempdir)
            first = folder / "任意名称A.pdf"
            second = folder / "完全不同的名字B.pdf"
            first.write_bytes(self.template_path.read_bytes())
            second.write_bytes(self.template_path.read_bytes())
            records = app.extract_all([first, second])
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0].source_file, first.name)

    def test_egg_contract_doubles_entry_price(self) -> None:
        egg_text = self.template_text.replace("LH2701", "JD2701")
        record = app.parse_confirmation(egg_text, "鸡蛋分支测试.pdf")
        self.assertTrue(record.is_egg)
        self.assertEqual(record.entry_price, Decimal("25010.00"))
        self.assertEqual(record.nominal_principal, Decimal("4801920.00"))

    def test_writes_second_data_row_with_formula_and_formats(self) -> None:
        source = HERE / "保期数据统计表.xlsx"
        with tempfile.TemporaryDirectory(prefix="confirmation_test_") as tempdir:
            output = Path(tempdir) / "result.xlsx"
            sheet, first_row, last_row = app.write_workbook(source, output, [self.record])
            self.assertEqual((sheet, first_row, last_row), ("Sheet1", 6, 6))
            ws = load_workbook(output, data_only=False)["Sheet1"]
            self.assertEqual(ws["U6"].value, "期货")
            self.assertEqual(ws["X6"].value, "=ROUND(V6*W6,2)")
            self.assertEqual(ws["AA6"].value, 1)
            self.assertEqual(ws["AA6"].number_format, "0.00%")
            self.assertEqual(ws["AC6"].number_format, "yyyy-mm-dd")
            self.assertEqual(ws["AN6"].value, "【HFSY】0147-JY-2026082701")
            self.assertEqual(ws["AN6"].number_format, "@")


if __name__ == "__main__":
    unittest.main()
