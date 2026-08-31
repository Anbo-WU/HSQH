#!/usr/bin/env python3
"""扫描商品交易确认书 PDF，并把结果写入统计表 U:AE 及 AN。

默认行为：
1. 扫描脚本所在目录中的 PDF，始终排除文件名含“模板确认书”的文件；
2. 自动定位表头行和现有第一条参考数据，从第二条数据行开始写入；
3. 不覆盖源统计表，生成“保期数据统计表_已填充.xlsx”；
4. 如果目标单元格已有内容，除非显式使用 --force，否则停止并提示。

常用命令：
    python3 扫描确认书填充数据统计表.py --preview
    python3 扫描确认书填充数据统计表.py
    python3 扫描确认书填充数据统计表.py --output 自定义结果.xlsx
    python3 扫描确认书填充数据统计表.py --self-test

依赖：openpyxl。PDF 文字提取依次尝试 pypdf、pdftotext；在 macOS 上还会
自动使用系统 PDFKit，并在 PDF 没有文字层时通过 Vision 做中文 OCR。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from copy import copy
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook


TWO_PLACES = Decimal("0.01")
TARGET_COLUMNS = {
    "U": "对冲方式",
    "V": "入场价格",
    "W": "名义数量",
    "X": "名义本金",
    "Y": "权利金",
    "Z": "场外期权型",
    "AA": "对冲比例",
    "AB": "对冲标的合约",
    "AC": "入场时间",
    "AD": "交易确认书生效日",
    "AE": "交易确认书到期日",
    "AN": "交易确认书编号",
}


MAC_PDF_EXTRACTOR_SOURCE = r'''
#import <Foundation/Foundation.h>
#import <AppKit/AppKit.h>
#import <PDFKit/PDFKit.h>
#import <Vision/Vision.h>
#import <math.h>

static NSString *OCRPage(PDFPage *page) {
    NSRect box = [page boundsForBox:kPDFDisplayBoxMediaBox];
    CGFloat scale = 2.5;
    NSImage *image = [page thumbnailOfSize:NSMakeSize(box.size.width * scale,
                                                       box.size.height * scale)
                                    forBox:kPDFDisplayBoxMediaBox];
    CGImageRef cgImage = [image CGImageForProposedRect:NULL context:nil hints:nil];
    if (!cgImage) return @"";

    VNRecognizeTextRequest *request = [[VNRecognizeTextRequest alloc] init];
    request.recognitionLevel = VNRequestTextRecognitionLevelAccurate;
    request.recognitionLanguages = @[@"zh-Hans", @"en-US"];
    request.usesLanguageCorrection = YES;

    VNImageRequestHandler *handler = [[VNImageRequestHandler alloc]
        initWithCGImage:cgImage options:@{}];
    NSError *error = nil;
    if (![handler performRequests:@[request] error:&error]) return @"";

    NSArray<VNRecognizedTextObservation *> *observations =
        [(NSArray<VNRecognizedTextObservation *> *)request.results
            sortedArrayUsingComparator:^NSComparisonResult(
                VNRecognizedTextObservation *a, VNRecognizedTextObservation *b) {
                CGFloat ay = a.boundingBox.origin.y + a.boundingBox.size.height;
                CGFloat by = b.boundingBox.origin.y + b.boundingBox.size.height;
                if (fabs(ay - by) > 0.015) {
                    return ay > by ? NSOrderedAscending : NSOrderedDescending;
                }
                if (a.boundingBox.origin.x < b.boundingBox.origin.x) return NSOrderedAscending;
                if (a.boundingBox.origin.x > b.boundingBox.origin.x) return NSOrderedDescending;
                return NSOrderedSame;
            }];

    NSMutableString *result = [NSMutableString string];
    for (VNRecognizedTextObservation *observation in observations) {
        VNRecognizedText *candidate = [[observation topCandidates:1] firstObject];
        if (candidate) [result appendFormat:@"%@\n", candidate.string];
    }
    return result;
}

int main(int argc, const char *argv[]) {
    @autoreleasepool {
        if (argc != 2) {
            fputs("usage: pdf_text PDF_PATH\n", stderr);
            return 2;
        }
        NSString *path = [NSString stringWithUTF8String:argv[1]];
        PDFDocument *document = [[PDFDocument alloc]
            initWithURL:[NSURL fileURLWithPath:path]];
        if (!document) {
            fputs("cannot open PDF\n", stderr);
            return 1;
        }
        for (NSInteger index = 0; index < document.pageCount; index++) {
            PDFPage *page = [document pageAtIndex:index];
            NSString *text = page.string ?: @"";
            if (text.length < 30) text = OCRPage(page);
            printf("%s\n\n---PAGE---\n\n", [text UTF8String]);
        }
    }
    return 0;
}
'''


@dataclass(frozen=True)
class ConfirmationData:
    source_file: str
    transaction_id: str
    hedge_method: str
    entry_price: Decimal
    nominal_quantity: Decimal
    nominal_principal: Decimal
    premium_total: Decimal
    otc_option_type: str
    hedge_ratio: Decimal
    contract: str
    signing_date: date
    effective_date: date
    expiry_date: date
    is_egg: bool

    def preview_dict(self) -> dict[str, str]:
        return {
            "文件": self.source_file,
            "交易编号": self.transaction_id,
            "U": self.hedge_method,
            "V": f"{self.entry_price:.2f}",
            "W": f"{self.nominal_quantity:.2f}",
            "X": f"{self.nominal_principal:.2f}",
            "Y": f"{self.premium_total:.2f}",
            "Z": self.otc_option_type,
            "AA": "100.00%",
            "AB": self.contract,
            "AC": self.signing_date.isoformat(),
            "AD": self.effective_date.isoformat(),
            "AE": self.expiry_date.isoformat(),
            "AN": self.transaction_id,
        }


class PDFTextExtractor:
    """在一次运行中复用 PDF 提取器；macOS helper 只编译一次。"""

    def __init__(self) -> None:
        self._tempdir: tempfile.TemporaryDirectory[str] | None = None
        self._mac_helper: Path | None = None

    def close(self) -> None:
        if self._tempdir is not None:
            self._tempdir.cleanup()
            self._tempdir = None
            self._mac_helper = None

    def __enter__(self) -> "PDFTextExtractor":
        return self

    def __exit__(self, *_: object) -> None:
        self.close()

    @staticmethod
    def _looks_usable(text: str) -> bool:
        return len(text.strip()) >= 80 and sum(
            marker in text for marker in ("交易确认书", "生效日", "标的合约", "入场价格")
        ) >= 2

    def extract(self, pdf_path: Path) -> str:
        errors: list[str] = []

        try:
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(str(pdf_path))
            text = "\n\n---PAGE---\n\n".join(page.extract_text() or "" for page in reader.pages)
            if self._looks_usable(text):
                return text
            errors.append("pypdf 未得到可用文字")
        except Exception as exc:  # pypdf 是可选依赖
            errors.append(f"pypdf: {exc}")

        pdftotext = shutil.which("pdftotext")
        if pdftotext:
            try:
                result = subprocess.run(
                    [pdftotext, "-layout", str(pdf_path), "-"],
                    check=True,
                    capture_output=True,
                    text=True,
                )
                if self._looks_usable(result.stdout):
                    return result.stdout
                errors.append("pdftotext 未得到可用文字")
            except Exception as exc:
                errors.append(f"pdftotext: {exc}")

        if platform.system() == "Darwin":
            try:
                helper = self._ensure_mac_helper()
                result = subprocess.run(
                    [str(helper), str(pdf_path)],
                    check=True,
                    capture_output=True,
                    text=True,
                )
                if self._looks_usable(result.stdout):
                    return result.stdout
                errors.append("macOS PDFKit/Vision 未得到可用文字")
            except Exception as exc:
                errors.append(f"macOS PDFKit/Vision: {exc}")

        details = "；".join(errors)
        raise RuntimeError(
            f"无法识别 PDF：{pdf_path.name}。{details}。"
            "非 macOS 扫描件请安装 pypdf，并为纯图片 PDF 安装可用的 OCR 工具。"
        )

    def _ensure_mac_helper(self) -> Path:
        if self._mac_helper is not None:
            return self._mac_helper
        clang = shutil.which("clang")
        if not clang:
            raise RuntimeError("未找到 clang")
        self._tempdir = tempfile.TemporaryDirectory(prefix="confirmation_pdf_")
        temp_root = Path(self._tempdir.name)
        source = temp_root / "pdf_text.m"
        binary = temp_root / "pdf_text"
        module_cache = temp_root / "module-cache"
        module_cache.mkdir()
        source.write_text(MAC_PDF_EXTRACTOR_SOURCE, encoding="utf-8")
        subprocess.run(
            [
                clang,
                "-fobjc-arc",
                "-fblocks",
                f"-fmodules-cache-path={module_cache}",
                "-framework",
                "Foundation",
                "-framework",
                "AppKit",
                "-framework",
                "PDFKit",
                "-framework",
                "Vision",
                str(source),
                "-o",
                str(binary),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        self._mac_helper = binary
        return binary


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\u00a0", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return text


def require_match(pattern: str, text: str, field_name: str, flags: int = 0) -> str:
    match = re.search(pattern, text, flags)
    if not match:
        raise ValueError(f"未找到字段：{field_name}")
    return match.group(1).strip()


def parse_decimal(raw: str, field_name: str) -> Decimal:
    try:
        return Decimal(raw.replace(",", "").strip()).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)
    except InvalidOperation as exc:
        raise ValueError(f"{field_name}不是有效数值：{raw!r}") from exc


def parse_date(raw: str, field_name: str) -> date:
    numbers = re.findall(r"\d+", raw)
    if len(numbers) < 3:
        raise ValueError(f"{field_name}不是有效日期：{raw!r}")
    try:
        return date(int(numbers[0]), int(numbers[1]), int(numbers[2]))
    except ValueError as exc:
        raise ValueError(f"{field_name}不是有效日期：{raw!r}") from exc


def clean_label_value(raw: str) -> str:
    return re.sub(r"[\s【】\[\]]+", "", raw).strip(" :：")


def parse_confirmation(text: str, source_file: str) -> ConfirmationData:
    text = normalize_text(text)
    section_mark = r"\s*[.．]\s*"
    date_pattern = r"(\d{4}\s*(?:年|[-/.])\s*\d{1,2}\s*(?:月|[-/.])\s*\d{1,2}\s*日?)"

    transaction_id = re.sub(
        r"\s+",
        "",
        require_match(r"交易编号\s*[:：]\s*([^\n]+)", text, "交易编号"),
    ).strip(" :：")
    signing_date = parse_date(
        require_match(r"签订时间\s*[:：]\s*" + date_pattern, text, "签订时间"),
        "签订时间",
    )
    effective_date = parse_date(
        require_match(r"2" + section_mark + r"1\s*[、,]?\s*生效日\s*[:：]\s*" + date_pattern,
                      text, "生效日"),
        "生效日",
    )
    expiry_date = parse_date(
        require_match(r"2" + section_mark + r"2\s*[、,]?\s*到期日\s*[:：]\s*" + date_pattern,
                      text, "到期日"),
        "到期日",
    )
    option_type = clean_label_value(
        require_match(r"3" + section_mark + r"1\s*[、,]?\s*期权类型\s*[:：]\s*([^\n]+)",
                      text, "期权类型")
    )
    option_structure = clean_label_value(
        require_match(r"3" + section_mark + r"2\s*[、,]?\s*期权结构\s*[:：]\s*([^\n]+)",
                      text, "期权结构")
    )
    contract = require_match(
        r"4" + section_mark + r"1\s*[、,]?\s*标的合约\s*[:：]\s*([A-Za-z]+\s*[-]?\s*\d{3,4})",
        text,
        "标的合约",
    )
    contract = re.sub(r"\s+", "", contract).upper()
    quantity = parse_decimal(
        require_match(
            r"4" + section_mark + r"1\s*[、,]?[\s\S]{0,160}?数量\s*[:：]\s*([0-9][0-9,]*(?:\.\d+)?)",
            text,
            "数量",
        ),
        "数量",
    )
    entry_price = parse_decimal(
        require_match(
            r"5" + section_mark + r"1\s*[、,]?\s*入场价格\s*[:：]\s*([0-9][0-9,]*(?:\.\d+)?)",
            text,
            "入场价格",
        ),
        "入场价格",
    )
    premium_total = parse_decimal(
        require_match(
            r"5" + section_mark + r"5\s*[、,]?\s*期权费\s*[:：][\s\S]{0,260}?"
            r"合计\s*[:：]\s*([0-9][0-9,]*(?:\.\d+)?)",
            text,
            "期权费合计",
        ),
        "期权费合计",
    )

    is_egg = "鸡蛋" in text or bool(re.match(r"^JD\d", contract, re.IGNORECASE))
    if is_egg:
        entry_price = (entry_price * Decimal("2")).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)
    principal = (entry_price * quantity).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)
    otc_option_type = f"{option_structure}{option_type}"

    if quantity <= 0 or entry_price <= 0 or premium_total < 0:
        raise ValueError("数量、入场价格必须为正数，权利金不能为负数")
    if not (signing_date <= effective_date <= expiry_date):
        raise ValueError(
            f"日期顺序异常：签订日 {signing_date}、生效日 {effective_date}、到期日 {expiry_date}"
        )

    return ConfirmationData(
        source_file=source_file,
        transaction_id=transaction_id,
        hedge_method="期货",
        entry_price=entry_price,
        nominal_quantity=quantity,
        nominal_principal=principal,
        premium_total=premium_total,
        otc_option_type=otc_option_type,
        hedge_ratio=Decimal("1.00"),
        contract=contract,
        signing_date=signing_date,
        effective_date=effective_date,
        expiry_date=expiry_date,
        is_egg=is_egg,
    )


def discover_pdfs(folder: Path) -> list[Path]:
    pdfs = [
        path
        for path in folder.iterdir()
        if path.is_file()
        and path.suffix.lower() == ".pdf"
        and "模板确认书" not in path.stem
    ]
    return sorted(pdfs, key=lambda path: path.name)


def extract_all(pdf_paths: Iterable[Path]) -> list[ConfirmationData]:
    results: list[ConfirmationData] = []
    failures: list[str] = []
    seen_content: dict[str, str] = {}
    with PDFTextExtractor() as extractor:
        for pdf_path in pdf_paths:
            try:
                text = extractor.extract(pdf_path)
                fingerprint_source = re.sub(r"\s+", "", normalize_text(text))
                fingerprint = hashlib.sha256(fingerprint_source.encode("utf-8")).hexdigest()
                if fingerprint in seen_content:
                    print(
                        f"跳过重复 PDF：{pdf_path.name}（内容与 {seen_content[fingerprint]} 相同）",
                        file=sys.stderr,
                    )
                    continue
                seen_content[fingerprint] = pdf_path.name
                results.append(parse_confirmation(text, pdf_path.name))
            except Exception as exc:
                failures.append(f"{pdf_path.name}: {exc}")
    if failures:
        raise RuntimeError("以下确认书识别失败：\n- " + "\n- ".join(failures))
    return results


def _cell_text(value: object) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def locate_rows(ws: object) -> tuple[int, int, int]:
    header_row = 0
    for row in range(1, min(getattr(ws, "max_row"), 30) + 1):
        if "对冲方式" in _cell_text(ws[f"U{row}"].value):
            header_row = row
            break
    if not header_row:
        raise ValueError("未在 U 列找到“对冲方式”表头，无法确认写入位置")

    for column, keyword in TARGET_COLUMNS.items():
        actual = _cell_text(ws[f"{column}{header_row}"].value)
        if keyword not in actual:
            raise ValueError(
                f"表头校验失败：{column}{header_row} 应包含“{keyword}”，实际为“{actual}”"
            )
    reference_row = header_row + 1
    start_row = header_row + 2
    return header_row, reference_row, start_row


def copy_reference_format(ws: object, reference_row: int, target_row: int) -> None:
    for column in TARGET_COLUMNS:
        source = ws[f"{column}{reference_row}"]
        target = ws[f"{column}{target_row}"]
        if source.has_style:
            target._style = copy(source._style)
    if ws.row_dimensions[reference_row].height is not None:
        ws.row_dimensions[target_row].height = ws.row_dimensions[reference_row].height


def write_workbook(
    workbook_path: Path,
    output_path: Path,
    records: Sequence[ConfirmationData],
    start_row_override: int | None = None,
    force: bool = False,
) -> tuple[str, int, int]:
    if workbook_path.suffix.lower() != ".xlsx":
        raise ValueError("当前统计表必须是 .xlsx；旧版 .xls 请先另存为 .xlsx")
    if output_path.exists() and output_path.resolve() != workbook_path.resolve() and not force:
        raise FileExistsError(f"输出文件已存在：{output_path}；确认覆盖请使用 --force")

    workbook = load_workbook(workbook_path)
    if len(workbook.sheetnames) != 1:
        raise ValueError(f"统计表应只有一个工作表，当前为：{workbook.sheetnames}")
    ws = workbook[workbook.sheetnames[0]]
    _, reference_row, detected_start_row = locate_rows(ws)
    start_row = start_row_override or detected_start_row
    if start_row <= reference_row:
        raise ValueError(f"起始行必须大于参考数据行 {reference_row}")

    target_rows = range(start_row, start_row + len(records))
    conflicts: list[str] = []
    for row in target_rows:
        occupied = [column for column in TARGET_COLUMNS if ws[f"{column}{row}"].value not in (None, "")]
        if occupied:
            conflicts.append(f"第 {row} 行（{','.join(occupied)}）")
    if conflicts and not force:
        raise ValueError(
            "目标区域已有数据：" + "；".join(conflicts) + "。确认覆盖请使用 --force"
        )

    for offset, record in enumerate(records):
        row = start_row + offset
        copy_reference_format(ws, reference_row, row)
        ws[f"U{row}"] = record.hedge_method
        ws[f"V{row}"] = float(record.entry_price)
        ws[f"W{row}"] = float(record.nominal_quantity)
        ws[f"X{row}"] = f"=ROUND(V{row}*W{row},2)"
        ws[f"Y{row}"] = float(record.premium_total)
        ws[f"Z{row}"] = record.otc_option_type
        ws[f"AA{row}"] = float(record.hedge_ratio)
        ws[f"AB{row}"] = record.contract
        ws[f"AC{row}"] = record.signing_date
        ws[f"AD{row}"] = record.effective_date
        ws[f"AE{row}"] = record.expiry_date
        ws[f"AN{row}"] = record.transaction_id

        for column in ("V", "W", "X", "Y"):
            ws[f"{column}{row}"].number_format = "0.00"
        for column in ("U", "Z", "AB", "AN"):
            ws[f"{column}{row}"].number_format = "@"
        ws[f"AA{row}"].number_format = "0.00%"
        for column in ("AC", "AD", "AE"):
            ws[f"{column}{row}"].number_format = "yyyy-mm-dd"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}_", suffix=".xlsx", dir=output_path.parent
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return ws.title, start_row, start_row + len(records) - 1


def print_preview(records: Sequence[ConfirmationData], as_json: bool = False) -> None:
    rows = [record.preview_dict() for record in records]
    if as_json:
        print(json.dumps(rows, ensure_ascii=False, indent=2))
        return
    columns = ["文件", "交易编号", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AN"]
    print("\t".join(columns))
    for row in rows:
        print("\t".join(row[column] for column in columns))


def self_test(folder: Path) -> None:
    template = folder / "模板确认书.pdf"
    if not template.exists():
        raise FileNotFoundError(f"未找到测试模板：{template}")
    with PDFTextExtractor() as extractor:
        template_text = extractor.extract(template)
        actual = parse_confirmation(template_text, template.name)
    expected = {
        "transaction_id": "【HFSY】0147-JY-2026082701",
        "entry_price": Decimal("12505.00"),
        "nominal_quantity": Decimal("192.00"),
        "nominal_principal": Decimal("2400960.00"),
        "premium_total": Decimal("101560.32"),
        "otc_option_type": "增强亚式看跌",
        "contract": "LH2701",
        "signing_date": date(2026, 8, 27),
        "effective_date": date(2026, 8, 28),
        "expiry_date": date(2026, 10, 28),
    }
    for field, expected_value in expected.items():
        actual_value = getattr(actual, field)
        if actual_value != expected_value:
            raise AssertionError(f"模板校验失败：{field}={actual_value!r}，应为 {expected_value!r}")
    if template in discover_pdfs(folder):
        raise AssertionError("生产文件发现逻辑没有排除模板确认书")
    egg = parse_confirmation(template_text.replace("LH2701", "JD2701"), "鸡蛋分支测试.pdf")
    if egg.entry_price != Decimal("25010.00") or egg.nominal_principal != Decimal("4801920.00"):
        raise AssertionError("鸡蛋入场价乘以 2 的规则校验失败")
    print("模板字段、鸡蛋价格乘 2 及生产扫描排除模板的校验均通过。")


def find_workbook(folder: Path, explicit: str | None) -> Path:
    if explicit:
        path = Path(explicit).expanduser()
        if not path.is_absolute():
            path = folder / path
        return path.resolve()
    candidates = sorted(
        path for path in folder.glob("*.xlsx") if "_已填充" not in path.stem and not path.name.startswith("~$")
    )
    if len(candidates) != 1:
        raise ValueError(f"应找到唯一的原始 .xlsx 统计表，实际为：{[p.name for p in candidates]}")
    return candidates[0].resolve()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="扫描确认书并填充保期数据统计表 U:AE 及 AN")
    parser.add_argument("--folder", default=None, help="数据目录；默认是脚本所在目录")
    parser.add_argument("--workbook", default=None, help="统计表路径；默认自动寻找唯一原始 .xlsx")
    parser.add_argument("--output", default=None, help="输出路径；默认生成 *_已填充.xlsx")
    parser.add_argument("--preview", action="store_true", help="只提取并打印，不写统计表")
    parser.add_argument("--json", action="store_true", help="预览时输出 JSON")
    parser.add_argument("--force", action="store_true", help="允许覆盖输出文件或目标区域已有数据")
    parser.add_argument("--start-row", type=int, default=None, help="覆盖自动定位的写入起始行")
    parser.add_argument("--self-test", action="store_true", help="仅使用模板确认书运行字段定位测试")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    folder = Path(args.folder).expanduser().resolve() if args.folder else Path(__file__).resolve().parent
    if args.self_test:
        self_test(folder)
        return 0

    pdfs = discover_pdfs(folder)
    if not pdfs:
        raise FileNotFoundError(f"{folder} 中没有找到真实确认书 PDF（模板确认书会被排除）")
    records = extract_all(pdfs)
    if args.preview:
        print_preview(records, as_json=args.json)
        return 0

    workbook_path = find_workbook(folder, args.workbook)
    if args.output:
        output_path = Path(args.output).expanduser()
        if not output_path.is_absolute():
            output_path = folder / output_path
        output_path = output_path.resolve()
    else:
        output_path = workbook_path.with_name(f"{workbook_path.stem}_已填充.xlsx")

    sheet_name, first_row, last_row = write_workbook(
        workbook_path,
        output_path,
        records,
        start_row_override=args.start_row,
        force=args.force,
    )
    print(f"完成：{len(records)} 份确认书已写入 {output_path}")
    print(
        f"工作表：{sheet_name}；范围：U{first_row}:AE{last_row}、AN{first_row}:AN{last_row}；"
        "模板确认书未参与生产扫描。"
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"错误：{exc}", file=sys.stderr)
        raise SystemExit(1)
