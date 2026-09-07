#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""根据评分标准计算每家客户的总分及风险承受能力等级。"""

from __future__ import annotations

import argparse
import os
import re
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

try:
    from openpyxl import load_workbook
except ImportError:
    print("错误：缺少 openpyxl。请先运行：python3 -m pip install openpyxl", file=sys.stderr)
    raise SystemExit(1)


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = BASE_DIR / "分数统计.xlsx"
QUESTION_COUNT = 20
MULTI_SELECT_QUESTIONS = {13, 17}

# 由“评分标准abv代码.txt”逐题翻译而来。
# 第13、17题虽是多选题，仍使用同一映射表，计算时取所选答案中的最高分。
SCORE_RULES: dict[int, dict[str, int]] = {
    1: {"A": 5, "B": 2, "C": 4, "D": 6},
    2: {"A": 1, "B": 2, "C": 3, "D": 5},
    3: {"A": 1, "B": 2, "C": 3, "D": 4},
    4: {"A": 1, "B": 2, "C": 3, "D": 4, "E": 0},
    5: {"A": 3, "B": 2, "C": 1, "D": 0, "E": 4},
    6: {"A": 1, "B": 3, "C": 4, "D": 5},
    7: {"A": 6, "B": 6, "C": 6, "D": 0},
    8: {"A": 1, "B": 3, "C": 5},
    9: {"A": 1, "B": 2, "C": 3, "D": 4},
    10: {"A": 0, "B": 1, "C": 2, "D": 4, "E": 5},
    11: {"A": 1, "B": 3, "C": 4, "D": 5},
    12: {"A": 1, "B": 2, "C": 3, "D": 4},
    13: {"A": 2, "B": 4, "C": 5, "D": 5, "E": 6, "F": 6},
    14: {"A": 1, "B": 2, "C": 3, "D": 4, "E": 0},
    15: {"A": 1, "B": 3, "C": 5},
    16: {"A": 0, "B": 2, "C": 4, "D": 5},
    17: {"A": 2, "B": 4, "C": 5, "D": 6, "E": 6, "F": 6},
    18: {"A": 0, "B": 2, "C": 4, "D": 6},
    19: {"A": 0, "B": 1, "C": 3, "D": 5, "E": 6},
    20: {"A": 3, "B": 5, "C": 4, "D": 1},
}

EXPECTED_HEADERS = ["客户名称", *range(1, QUESTION_COUNT + 1), "总分计算", "风险承受能力等级"]


class ScoringError(Exception):
    """表格或答案不符合评分要求。"""


@dataclass(frozen=True)
class ScoredCustomer:
    row_number: int
    customer_name: str
    total_score: int
    risk_level: str


def normalize_answer(value: object, question_number: int) -> str:
    """规范答案，并严格检查单选、多选及该题允许的选项。"""
    if value is None:
        raise ScoringError(f"第 {question_number} 题答案为空")

    raw_answer = str(value).strip().upper()
    if not raw_answer:
        raise ScoringError(f"第 {question_number} 题答案为空")

    # 多选题兼容 ABC、A B、A、B、A/B、A和B 等写法。
    letters = re.findall(r"[A-Z]", raw_answer)
    residue = re.sub(r"[A-Z\s,，、;/／+和]", "", raw_answer)
    if residue or not letters:
        raise ScoringError(f"第 {question_number} 题答案格式无效：{value}")

    if len(set(letters)) != len(letters):
        raise ScoringError(f"第 {question_number} 题含重复选项：{''.join(letters)}")

    if question_number not in MULTI_SELECT_QUESTIONS and len(letters) != 1:
        raise ScoringError(
            f"第 {question_number} 题是单选题，但检测到多个答案：{''.join(letters)}"
        )

    invalid = [letter for letter in letters if letter not in SCORE_RULES[question_number]]
    if invalid:
        allowed = "、".join(SCORE_RULES[question_number])
        raise ScoringError(
            f"第 {question_number} 题存在无效选项“{'、'.join(invalid)}”，允许的选项为 {allowed}"
        )
    return "".join(letters)


def score_answer(question_number: int, value: object) -> int:
    answer = normalize_answer(value, question_number)
    scores = [SCORE_RULES[question_number][letter] for letter in answer]
    if question_number in MULTI_SELECT_QUESTIONS:
        return max(scores)
    return scores[0]


def calculate_total(answers: Iterable[object]) -> int:
    answer_list = list(answers)
    if len(answer_list) != QUESTION_COUNT:
        raise ScoringError(
            f"计算总分需要 {QUESTION_COUNT} 个答案，实际收到 {len(answer_list)} 个"
        )
    return sum(
        score_answer(question_number, answer)
        for question_number, answer in enumerate(answer_list, start=1)
    )


def risk_level(total_score: int) -> str:
    if not 0 <= total_score <= 100:
        raise ScoringError(f"总分超出有效范围 0～100：{total_score}")
    if total_score < 20:
        return "C1 保守型"
    if total_score <= 30:
        return "C2 谨慎型"
    if total_score <= 50:
        return "C3 稳健型"
    if total_score <= 85:
        return "C4 积极型"
    return "C5 激进型"


def normalize_header(value: object) -> object:
    """允许题号表头为数字 1 或文本“1”，其余表头保持去空格后的文本。"""
    if isinstance(value, str):
        value = value.strip()
        if value.isdigit():
            return int(value)
    return value


def validate_headers(worksheet) -> None:
    actual = [normalize_header(worksheet.cell(1, column).value) for column in range(1, 24)]
    if actual != EXPECTED_HEADERS:
        raise ScoringError(
            "Excel 的 A1:W1 表头与预期不一致；应为“客户名称、1～20、总分计算、风险承受能力等级”"
        )


def collect_scores(worksheet) -> list[ScoredCustomer]:
    """先检查并计算全部客户；任一行出错时不进入写入阶段。"""
    records: list[ScoredCustomer] = []
    errors: list[str] = []

    for row_number in range(2, worksheet.max_row + 1):
        customer_name_value = worksheet.cell(row_number, 1).value
        answers = [worksheet.cell(row_number, column).value for column in range(2, 22)]

        # 完全空白的 A:U 行不是客户数据，直接跳过。
        if customer_name_value is None and all(value is None for value in answers):
            continue

        customer_name = "" if customer_name_value is None else str(customer_name_value).strip()
        if not customer_name:
            errors.append(f"第 {row_number} 行：客户名称为空")
            continue

        try:
            total = calculate_total(answers)
        except ScoringError as exc:
            errors.append(f"第 {row_number} 行（{customer_name}）：{exc}")
            continue

        records.append(
            ScoredCustomer(
                row_number=row_number,
                customer_name=customer_name,
                total_score=total,
                risk_level=risk_level(total),
            )
        )

    if errors:
        raise ScoringError("以下数据无法评分：\n  - " + "\n  - ".join(errors))
    if not records:
        raise ScoringError("Excel 中没有找到可评分的客户数据")
    return records


def save_atomically(workbook, workbook_path: Path) -> None:
    """先保存为同目录临时文件，成功后再替换原表，降低文件损坏风险。"""
    workbook_path = workbook_path.resolve()
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{workbook_path.stem}_",
            suffix=workbook_path.suffix,
            dir=workbook_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
    except Exception as exc:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise ScoringError(
            f"保存 Excel 失败（请确认文件未被 Excel/WPS 占用）：{exc}"
        ) from exc


def update_workbook(workbook_path: Path, preview: bool = False) -> list[ScoredCustomer]:
    if not workbook_path.is_file():
        raise ScoringError(f"找不到 Excel 文件：{workbook_path}")

    try:
        workbook = load_workbook(workbook_path)
    except Exception as exc:
        raise ScoringError(f"无法打开 Excel 文件：{exc}") from exc

    worksheet = workbook.active
    validate_headers(worksheet)
    records = collect_scores(worksheet)

    if not preview:
        # 按表头位置写入：V 列为总分，W 列为风险承受能力等级。
        for record in records:
            worksheet.cell(record.row_number, 22, record.total_score)
            worksheet.cell(record.row_number, 23, record.risk_level)
        save_atomically(workbook, workbook_path)

    return records


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="读取分数统计.xlsx 中每家客户的20题答案，计算总分和风险等级。"
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="需要计算的 Excel 文件路径",
    )
    parser.add_argument("--preview", action="store_true", help="只显示计算结果，不修改 Excel")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        records = update_workbook(args.workbook, preview=args.preview)
    except ScoringError as exc:
        print(f"计算失败，Excel 未作任何修改：\n{exc}", file=sys.stderr)
        return 1

    print(f"已成功计算 {len(records)} 家客户：")
    for record in records:
        print(
            f"  第 {record.row_number} 行 | {record.customer_name} | "
            f"总分 {record.total_score} | {record.risk_level}"
        )

    if args.preview:
        print("当前为预览模式，未修改 Excel。")
    else:
        print(f"已写入：{args.workbook.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
